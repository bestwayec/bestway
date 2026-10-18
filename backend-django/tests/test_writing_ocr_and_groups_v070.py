"""Tests for v0.7.0 Writing Vision OCR, 3-Tier Gemini Fallback, Teacher Groups/Analytics, Excel Export, and Superadmin."""
from __future__ import annotations

import json
import os
from datetime import UTC, datetime
from io import BytesIO
from tempfile import TemporaryDirectory
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

from django.core.files.uploadedfile import SimpleUploadedFile
from django.core.management import call_command
from django.test import SimpleTestCase
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.core import assessment_providers, assessment_worker, writing_ocr
from tests.test_assessment_quality import fixture as quality_fixture


class WritingOcrAndGeminiFallbackTests(SimpleTestCase):
    def setUp(self):
        self.client = APIClient()

    def test_extract_and_evaluate_images_multi_page_with_gemini_fallback(self):
        page1 = (b"Page 1 essay: Education plays a pivotal role in modern society.", "image/jpeg")
        page2 = (b"Page 2 essay: Furthermore, technology empowers students globally.", "image/png")

        def fail_post(url, payload, headers, timeout=8.0):
            if "openrouter.ai" in url or "xkiro.com" in url:
                raise RuntimeError("Primary/secondary unavailable")
            if "generativelanguage.googleapis.com" in url:
                if "inline_data" in json.dumps(payload):
                    return {
                        "candidates": [
                            {
                                "content": {
                                    "parts": [
                                        {"text": "Extracted handwritten paragraph from Gemini Vision."}
                                    ]
                                }
                            }
                        ]
                    }
                return {
                    "candidates": [
                        {
                            "content": {
                                "parts": [
                                    {
                                        "text": json.dumps({
                                            "task_achievement": 7.0,
                                            "coherence_cohesion": 7.5,
                                            "lexical_resource": 7.0,
                                            "grammatical_range": 6.5,
                                            "overall_band": 7.0,
                                            "strengths": ["Strong paragraphing", "Academic vocabulary"],
                                            "weaknesses": ["Minor article omissions"],
                                            "feedback": "Well-structured Task 2 essay with clear arguments.",
                                        })
                                    }
                                ]
                            }
                        }
                    ]
                }
            raise RuntimeError("Unknown URL")

        with patch.dict(os.environ, {"OPENROUTER_API_KEY": "or-key", "XKIRO_API_KEY": "xk-key"}, clear=False), \
             patch.object(writing_ocr, "_post_json", side_effect=fail_post):
            res = writing_ocr.extract_and_evaluate_images([page1, page2], task_prompt="Discuss education.", task_type="task2")
            self.assertEqual(len(res["pages"]), 2)
            self.assertIn("Extracted handwritten paragraph from Gemini Vision.", res["ocrText"])
            self.assertEqual(res["overall_band"], 7.0)
            self.assertEqual(res["task_achievement"], 7.0)
            self.assertEqual(res["coherence_cohesion"], 7.5)
            self.assertEqual(res["provider"], "gemini")

    def test_assessment_providers_3_tier_chain_falls_back_to_gemini(self):
        input_data, valid_result = quality_fixture("writing")
        gemini_resp = {
            "candidates": [
                {
                    "content": {
                        "parts": [{"text": json.dumps(valid_result)}]
                    }
                }
            ],
            "usageMetadata": {"promptTokenCount": 140, "candidatesTokenCount": 65},
        }
        env_cfg = {
            "ASSESSMENT_PROVIDER": "openrouter",
            "OPENROUTER_API_KEY": "or-test-key",
            "XKIRO_API_KEY": "xk-test-key",
            "GEMINI_API_KEY": "gemini-test-key-v070",
        }

        def fake_http(url, body, headers, timeout):
            if "openrouter.ai" in url:
                raise assessment_providers.ProviderError("PROVIDER_RATE_LIMITED", True, False)
            if "xkiro.com" in url:
                raise assessment_providers.ProviderError("PROVIDER_UNAVAILABLE", True, False)
            if "generativelanguage.googleapis.com" in url:
                return gemini_resp
            raise AssertionError(f"Unexpected URL: {url}")

        with patch.dict(os.environ, env_cfg, clear=True), \
             patch.object(assessment_providers, "provider_json", side_effect=fake_http):
            output = assessment_providers.assess(input_data, "PRIMARY")
            self.assertEqual(output["provider"], "gemini")
            self.assertEqual(output["model"], "gemini-2.0-flash")
            self.assertEqual(output["inputTokens"], 140)
            self.assertEqual(output["outputTokens"], 65)

    def test_writing_ocr_upload_endpoint_saves_images_and_updates_mock_answer(self):
        student = SimpleNamespace(id="student-1", role="student", is_authenticated=True)
        self.client.force_authenticate(student)

        attempt = SimpleNamespace(id="att-100", student_id="student-1", exam_id="exam-1", status="in_progress")
        img1 = SimpleUploadedFile("page1.jpg", b"\xff\xd8\xff\xe0Handwritten page one text here.", content_type="image/jpeg")
        img2 = SimpleUploadedFile("page2.png", b"\x89PNG\r\n\x1a\nHandwritten page two text here.", content_type="image/png")

        with TemporaryDirectory() as tmp_dir, \
             patch.dict(os.environ, {"STORAGE_DIR": tmp_dir}), \
             patch("apps.legacy_schema.models.MockAttempt.objects.filter") as mock_att_filter, \
             patch("apps.legacy_schema.models.MockAnswer.objects.filter") as mock_ans_filter, \
             patch("apps.legacy_schema.models.MockAnswer.objects.update_or_create") as mock_upsert, \
             patch.object(writing_ocr, "extract_text_from_images", return_value={
                 "ocrText": "First page paragraph.\n\nSecond page conclusion.",
                 "pages": ["First page paragraph.", "Second page conclusion."],
                 "wordCount": 6,
                 "provider": "gemini",
             }):
            mock_att_filter.return_value.first.return_value = attempt
            mock_ans_filter.return_value.first.return_value = None

            response = self.client.post(
                "/v1/mock/attempts/att-100/writing/q-writing-2/ocr",
                {"images": [img1, img2]},
                format="multipart",
            )
            self.assertEqual(response.status_code, 200)
            data = response.data["data"]
            self.assertEqual(data["questionId"], "q-writing-2")
            self.assertEqual(data["ocrText"], "First page paragraph.\n\nSecond page conclusion.")
            self.assertEqual(data["wordCount"], 6)
            self.assertEqual(len(data["imageKeys"]), 2)
            self.assertEqual(len(data["imageUrls"]), 2)
            mock_upsert.assert_called_once()
            call_kwargs = mock_upsert.call_args.kwargs
            self.assertEqual(call_kwargs["attempt_id"], "att-100")
            self.assertEqual(call_kwargs["question_id"], "q-writing-2")
            self.assertEqual(call_kwargs["defaults"]["response"], "First page paragraph.\n\nSecond page conclusion.")
            self.assertEqual(len(call_kwargs["defaults"]["rubric_scores"]["imageKeys"]), 2)

    def test_assessment_worker_finish_triggers_telegram_notification(self):
        fake_job = SimpleNamespace(id="job-99", attempt_id="att-99", skill="writing", status="SUCCEEDED")
        with patch("apps.core.telegram_delivery.notify_submission_result", create=True) as mock_notify:
            assessment_worker.finish(fake_job)
            mock_notify.assert_called_once_with(fake_job)


class TeacherGroupsAnalyticsAndExcelExportTests(SimpleTestCase):
    def setUp(self):
        self.client = APIClient()

    def test_teacher_can_create_group_and_defaults_teacher_id(self):
        teacher = SimpleNamespace(id="teacher-77", role="teacher", is_authenticated=True, name="Mr. Teacher")
        self.client.force_authenticate(teacher)

        created_group = SimpleNamespace(
            id="grp-new",
            name="IELTS Intensive 7.5",
            teacher_id="teacher-77",
            teacher=SimpleNamespace(name="Mr. Teacher"),
            schedule=[{"day": "mon", "startTime": "14:00", "endTime": "16:00"}],
            created_at=datetime(2026, 10, 16, tzinfo=UTC),
        )

        class DummyAtomic:
            def __enter__(self):
                return self
            def __exit__(self, *args):
                return False

        with patch("apps.core.groups.transaction.atomic", return_value=DummyAtomic()), \
             patch("apps.core.groups.Group.objects.create", return_value=created_group) as mock_create, \
             patch("apps.core.groups.audit"), \
             patch("apps.core.groups.get_group", return_value=created_group), \
             patch("apps.core.groups.StudentProfile.objects.filter") as mock_sp_filter:
            mock_sp_filter.return_value.select_related.return_value.order_by.return_value = []

            res = self.client.post(
                "/v1/groups",
                {
                    "name": "IELTS Intensive 7.5",
                    "schedule": [{"day": "mon", "startTime": "14:00", "endTime": "16:00"}],
                },
                format="json",
            )
            self.assertEqual(res.status_code, 201)
            self.assertEqual(res.data["data"]["teacherId"], "teacher-77")
            self.assertEqual(res.data["data"]["name"], "IELTS Intensive 7.5")
            mock_create.assert_called_once()
            self.assertEqual(mock_create.call_args.kwargs["teacher_id"], "teacher-77")

    def test_group_detail_includes_student_scores_and_group_analytics(self):
        teacher = SimpleNamespace(id="teacher-1", role="teacher", is_authenticated=True)
        self.client.force_authenticate(teacher)

        group_obj = SimpleNamespace(
            id="grp-1",
            name="IELTS Pro",
            teacher_id="teacher-1",
            teacher=SimpleNamespace(name="Teacher One"),
            schedule=[],
            created_at=datetime(2026, 10, 1, tzinfo=UTC),
        )
        student_profile = SimpleNamespace(
            user_id="stu-1",
            is_approved=True,
            current_points=180,
            user=SimpleNamespace(
                id="stu-1",
                name="Madina Karimova",
                phone="+998901112233",
                email="madina@example.com",
                is_active=True,
            ),
        )
        attempt_obj = SimpleNamespace(
            id="att-1",
            student_id="stu-1",
            status="completed",
            section_bands={"listening": 7.5, "reading": 7.0, "writing": 6.5, "speaking": 7.0},
            standard_scores=None,
            raw_scores=None,
            overall_band=7.0,
            overall_score=None,
            cefr_level="C1",
            finished_at=datetime(2026, 10, 15, 12, 0, tzinfo=UTC),
            submitted_at=datetime(2026, 10, 15, 12, 0, tzinfo=UTC),
            started_at=datetime(2026, 10, 15, 10, 0, tzinfo=UTC),
            exam=SimpleNamespace(title="IELTS Mock #1"),
        )

        with patch("apps.core.groups.get_group", return_value=group_obj), \
             patch("apps.core.groups.StudentProfile.objects.filter") as mock_sp, \
             patch("apps.core.groups.MockAttempt.objects.filter") as mock_att, \
             patch("apps.core.groups.TestAttempt.objects.filter") as mock_leg:
            mock_sp.return_value.select_related.return_value.order_by.return_value = [student_profile]
            mock_att.return_value.select_related.return_value.order_by.return_value = [attempt_obj]
            mock_leg.return_value.order_by.return_value = []

            res = self.client.get("/v1/groups/grp-1")
            self.assertEqual(res.status_code, 200)
            payload = res.data["data"]
            self.assertEqual(payload["analytics"]["averageBandScore"], 7.0)
            self.assertEqual(payload["analytics"]["totalCompletedAttempts"], 1)
            self.assertEqual(payload["analytics"]["activeStudentsCount"], 1)
            stu = payload["students"][0]
            self.assertEqual(stu["completedExamsCount"], 1)
            self.assertEqual(stu["writingScore"], 6.5)
            self.assertEqual(stu["speakingScore"], 7.0)
            self.assertEqual(stu["overallBand"], 7.0)

    def test_group_excel_export_generates_valid_openpyxl_workbook(self):
        teacher = SimpleNamespace(id="teacher-1", role="teacher", is_authenticated=True)
        self.client.force_authenticate(teacher)

        group_obj = SimpleNamespace(
            id="grp-excel",
            name="IELTS Band 8",
            teacher_id="teacher-1",
            teacher=SimpleNamespace(name="Teacher One"),
        )
        student_profile = SimpleNamespace(
            user_id="stu-1",
            user=SimpleNamespace(name="Sardor Aliyev", phone="+998901234567", email="sardor@bestway.uz"),
        )
        attempt_obj = SimpleNamespace(
            id="att-excel-1",
            student_id="stu-1",
            status="completed",
            section_bands={"listening": 8.0, "reading": 7.5, "writing": 7.0, "speaking": 7.5},
            overall_band=7.5,
            overall_score=None,
            cefr_level="C1",
            finished_at=datetime(2026, 10, 16, 9, 30, tzinfo=UTC),
            submitted_at=None,
            started_at=datetime(2026, 10, 16, 8, 0, tzinfo=UTC),
            exam=SimpleNamespace(title="Official IELTS Mock 4"),
        )

        with patch("apps.core.groups.get_group", return_value=group_obj), \
             patch("apps.core.groups.StudentProfile.objects.filter") as mock_sp, \
             patch("apps.core.groups.MockAttempt.objects.filter") as mock_att, \
             patch("apps.core.groups.MockAnswer.objects.filter") as mock_ans:
            mock_sp.return_value.select_related.return_value.order_by.return_value = [student_profile]
            mock_att.return_value.select_related.return_value.order_by.return_value = [attempt_obj]
            mock_ans.return_value.exclude.return_value.exclude.return_value = [
                SimpleNamespace(attempt_id="att-excel-1", feedback="Strong lexical resource and clear cohesion.")
            ]

            res = self.client.get("/v1/groups/grp-excel/export/excel")
            self.assertEqual(res.status_code, 200)
            self.assertEqual(
                res["Content-Type"],
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            )
            self.assertIn(".xlsx", res["Content-Disposition"])

            wb = load_workbook(BytesIO(res.content))
            ws = wb.active
            headers = [ws.cell(row=1, column=col).value for col in range(1, 13)]
            self.assertEqual(
                headers,
                [
                    "Student Name",
                    "Email / Phone",
                    "Group Name",
                    "Exam Type",
                    "Listening",
                    "Reading",
                    "Writing",
                    "Speaking",
                    "Overall Band",
                    "CEFR Level",
                    "Submitted Date",
                    "Teacher Feedback Summary",
                ],
            )
            self.assertEqual(ws.cell(row=2, column=1).value, "Sardor Aliyev")
            self.assertEqual(ws.cell(row=2, column=9).value, 7.5)
            self.assertEqual(ws.cell(row=3, column=1).value, "GROUP AVERAGE")

    def test_teacher_can_delete_mock_exam(self):
        teacher = SimpleNamespace(id="teacher-del", role="teacher", is_authenticated=True)
        self.client.force_authenticate(teacher)

        fake_exam = SimpleNamespace(id="exam-del-1", title="Draft Exam", delete=MagicMock())
        with patch("apps.legacy_schema.models.MockExam.objects.filter") as mock_filter, \
             patch("apps.core.mock_authoring.delete_exam", return_value={"deleted": True}):
            mock_filter.return_value.first.return_value = fake_exam
            res = self.client.delete("/v1/mock/exams/exam-del-1")
            self.assertEqual(res.status_code, 200)
            self.assertTrue(res.data["data"]["deleted"])
            self.assertEqual(res.data["data"]["id"], "exam-del-1")

    def test_ensure_superadmin_management_command_sets_telegram_id_5582178765(self):
        existing_admin = SimpleNamespace(
            id="super-1",
            email="superadmin@bestwayec.uz",
            role="admin",
            is_active=False,
            telegram_chat_id=None,
        )
        with patch("apps.core.management.commands.ensure_superadmin.User.objects.filter") as mock_user_filter, \
             patch("apps.core.management.commands.ensure_superadmin.StudentProfile.objects.filter"):
            mock_user_filter.return_value.order_by.return_value.first.return_value = existing_admin
            mock_update = mock_user_filter.return_value.update

            from io import StringIO
            out = StringIO()
            call_command("ensure_superadmin", "--telegram-id", "5582178765", stdout=out)
            mock_update.assert_called_once()
            update_kwargs = mock_update.call_args.kwargs
            self.assertEqual(update_kwargs["role"], "super_admin")
            self.assertTrue(update_kwargs["is_active"])
            self.assertEqual(update_kwargs["telegram_chat_id"], "5582178765")
