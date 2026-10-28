"""Teacher PIN-Code & Live Exam Room Service (Real-Time Monitoring + Excel Results Export)."""
from __future__ import annotations

import random
import re
from datetime import timedelta
from io import BytesIO
from uuid import uuid4

from django.db import connection
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.legacy_schema.models import (
    Group,
    MockAnswer,
    MockAttempt,
    MockExam,
    MockQuestion,
    StudentProfile,
    User,
)
from common.api.exceptions import ContractAPIException

from . import mock_attempts, mock_catalog
from .prediction_service import extract_attempt_dual_metrics

_live_tables_ensured = False


def ensure_live_exam_tables() -> None:
    global _live_tables_ensured
    if _live_tables_ensured:
        return
    with connection.cursor() as cursor:
        if connection.vendor == "postgresql":
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS "LiveExamSession" (
                    "id" VARCHAR(255) PRIMARY KEY,
                    "pin" VARCHAR(12) NOT NULL,
                    "exam_id" VARCHAR(255) NOT NULL,
                    "group_id" VARCHAR(255),
                    "teacher_id" VARCHAR(255) NOT NULL,
                    "title" TEXT,
                    "simulator_mode" VARCHAR(20) DEFAULT 'cdi',
                    "status" VARCHAR(20) DEFAULT 'active',
                    "duration_minutes" INTEGER DEFAULT 60,
                    "created_at" TIMESTAMP WITH TIME ZONE DEFAULT NOW(),
                    "ends_at" TIMESTAMP WITH TIME ZONE,
                    "closed_at" TIMESTAMP WITH TIME ZONE
                );
                CREATE INDEX IF NOT EXISTS "idx_live_exam_session_pin" ON "LiveExamSession" ("pin", "status");
                CREATE INDEX IF NOT EXISTS "idx_live_exam_session_teacher" ON "LiveExamSession" ("teacher_id", "created_at" DESC);

                CREATE TABLE IF NOT EXISTS "LiveExamParticipant" (
                    "id" VARCHAR(255) PRIMARY KEY,
                    "session_id" VARCHAR(255) NOT NULL,
                    "student_id" VARCHAR(255) NOT NULL,
                    "attempt_id" VARCHAR(255),
                    "status" VARCHAR(20) DEFAULT 'in_progress',
                    "current_skill" VARCHAR(30),
                    "current_question_number" INTEGER DEFAULT 1,
                    "answered_count" INTEGER DEFAULT 0,
                    "flagged_count" INTEGER DEFAULT 0,
                    "total_questions" INTEGER DEFAULT 40,
                    "anti_cheat_count" INTEGER DEFAULT 0,
                    "raw_score" FLOAT,
                    "band_score" FLOAT,
                    "multilevel_score" FLOAT,
                    "cefr_level" VARCHAR(10),
                    "joined_at" TIMESTAMP WITH TIME ZONE DEFAULT NOW(),
                    "finished_at" TIMESTAMP WITH TIME ZONE,
                    "updated_at" TIMESTAMP WITH TIME ZONE DEFAULT NOW()
                );
                CREATE UNIQUE INDEX IF NOT EXISTS "idx_live_participant_unique" ON "LiveExamParticipant" ("session_id", "student_id");

                ALTER TABLE "LiveExamParticipant" ADD COLUMN IF NOT EXISTS "time_remaining_seconds" INTEGER;
                ALTER TABLE "LiveExamParticipant" ADD COLUMN IF NOT EXISTS "strikes_count" INTEGER DEFAULT 0;
                ALTER TABLE "LiveExamParticipant" ADD COLUMN IF NOT EXISTS "last_heartbeat" TIMESTAMP WITH TIME ZONE DEFAULT NOW();
                ALTER TABLE "LiveExamParticipant" ADD COLUMN IF NOT EXISTS "is_focused" BOOLEAN DEFAULT TRUE;
                ALTER TABLE "LiveExamSession" ADD COLUMN IF NOT EXISTS "exam_state" VARCHAR(20) DEFAULT 'waiting';
                """
            )
        else:
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS "LiveExamSession" (
                    "id" VARCHAR(255) PRIMARY KEY,
                    "pin" VARCHAR(12) NOT NULL,
                    "exam_id" VARCHAR(255) NOT NULL,
                    "group_id" VARCHAR(255),
                    "teacher_id" VARCHAR(255) NOT NULL,
                    "title" TEXT,
                    "simulator_mode" VARCHAR(20) DEFAULT 'cdi',
                    "status" VARCHAR(20) DEFAULT 'active',
                    "duration_minutes" INTEGER DEFAULT 60,
                    "created_at" TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    "ends_at" TIMESTAMP,
                    "closed_at" TIMESTAMP
                );
                """
            )
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS "LiveExamParticipant" (
                    "id" VARCHAR(255) PRIMARY KEY,
                    "session_id" VARCHAR(255) NOT NULL,
                    "student_id" VARCHAR(255) NOT NULL,
                    "attempt_id" VARCHAR(255),
                    "status" VARCHAR(20) DEFAULT 'in_progress',
                    "current_skill" VARCHAR(30),
                    "current_question_number" INTEGER DEFAULT 1,
                    "answered_count" INTEGER DEFAULT 0,
                    "flagged_count" INTEGER DEFAULT 0,
                    "total_questions" INTEGER DEFAULT 40,
                    "anti_cheat_count" INTEGER DEFAULT 0,
                    "raw_score" FLOAT,
                    "band_score" FLOAT,
                    "multilevel_score" FLOAT,
                    "cefr_level" VARCHAR(10),
                    "joined_at" TIMESTAMP DEFAULT CURRENT_TIMESTAMP,
                    "finished_at" TIMESTAMP,
                    "updated_at" TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                );
                """
            )
            cursor.execute('PRAGMA table_info("LiveExamParticipant")')
            p_cols = {row[1] for row in cursor.fetchall()}
            if "time_remaining_seconds" not in p_cols:
                cursor.execute('ALTER TABLE "LiveExamParticipant" ADD COLUMN "time_remaining_seconds" INTEGER')
            if "strikes_count" not in p_cols:
                cursor.execute('ALTER TABLE "LiveExamParticipant" ADD COLUMN "strikes_count" INTEGER DEFAULT 0')
            if "last_heartbeat" not in p_cols:
                cursor.execute('ALTER TABLE "LiveExamParticipant" ADD COLUMN "last_heartbeat" TIMESTAMP DEFAULT CURRENT_TIMESTAMP')
            if "is_focused" not in p_cols:
                cursor.execute('ALTER TABLE "LiveExamParticipant" ADD COLUMN "is_focused" BOOLEAN DEFAULT 1')

            cursor.execute('PRAGMA table_info("LiveExamSession")')
            s_cols = {row[1] for row in cursor.fetchall()}
            if "exam_state" not in s_cols:
                cursor.execute('ALTER TABLE "LiveExamSession" ADD COLUMN "exam_state" VARCHAR(20) DEFAULT \'waiting\'')
    _live_tables_ensured = True


def _iso(val) -> str | None:
    if val is None:
        return None
    if hasattr(val, "isoformat"):
        return val.isoformat()
    return str(val)


def _generate_unique_pin() -> str:
    ensure_live_exam_tables()
    with connection.cursor() as cursor:
        for _ in range(30):
            pin = str(random.randint(100000, 999999))
            cursor.execute(
                'SELECT 1 FROM "LiveExamSession" WHERE "pin"=%s AND "status" IN (\'waiting\', \'active\')',
                [pin],
            )
            if not cursor.fetchone():
                return pin
    return str(random.randint(100000, 999999))


def _session_row_to_dict(row) -> dict:
    session_id, pin, exam_id, group_id, teacher_id, title, sim_mode, status, dur_min, created_at, ends_at, closed_at = row
    exam = MockExam.objects.filter(id=exam_id).first()
    group = Group.objects.filter(id=group_id).first() if group_id else None
    teacher = User.objects.filter(id=teacher_id).first()

    with connection.cursor() as cursor:
        cursor.execute('SELECT COUNT(*) FROM "LiveExamParticipant" WHERE "session_id"=%s', [session_id])
        part_count = cursor.fetchone()[0]

    return {
        "id": session_id,
        "roomId": session_id,
        "pin": pin,
        "joinUrl": f"/mock/live?pin={pin}",
        "examId": exam_id,
        "examTitle": exam.title if exam else (title or "Mock Exam"),
        "examType": exam.type if exam else "ielts_academic",
        "groupId": group_id,
        "groupName": group.name if group else None,
        "teacherId": teacher_id,
        "teacherName": teacher.name if teacher else None,
        "title": title or (exam.title if exam else "Live Exam Room"),
        "simulatorMode": sim_mode or ("dtm" if (exam and exam.type == "multilevel") else "cdi"),
        "status": status,
        "durationMinutes": int(dur_min or 60),
        "participantsCount": int(part_count or 0),
        "createdAt": _iso(created_at),
        "endsAt": _iso(ends_at),
        "closedAt": _iso(closed_at),
    }


def create_session(actor, data: dict) -> dict:
    """Create a new Live Exam Session with a unique 6-digit PIN."""
    ensure_live_exam_tables()
    if not actor or getattr(actor, "role", None) not in ("teacher", "admin", "super_admin"):
        raise ContractAPIException("FORBIDDEN", "Faqat o'qituvchi yoki admin jonli imtihon xonasini ocha oladi", 403)

    exam_id = (data.get("examId") or data.get("exam_id") or "").strip()
    if not exam_id:
        raise ContractAPIException("VALIDATION_ERROR", "examId ko'rsatilishi shart", 400)

    exam = MockExam.objects.filter(id=exam_id).first()
    if not exam:
        raise ContractAPIException("MOCK_EXAM_NOT_FOUND", "Mock imtihon topilmadi", 404)

    group_id = (data.get("groupId") or data.get("group_id") or "").strip() or None
    if group_id:
        group = Group.objects.filter(id=group_id).first()
        if not group:
            raise ContractAPIException("GROUP_NOT_FOUND", "Guruh topilmadi", 404)

    default_sim = "dtm" if exam.type == "multilevel" else "cdi"
    sim_mode = str(data.get("simulatorMode") or data.get("mode") or default_sim).strip().lower()
    if sim_mode not in ("cdi", "dtm"):
        sim_mode = default_sim

    try:
        duration_minutes = max(5, min(300, int(data.get("durationMinutes") or 60)))
    except (TypeError, ValueError):
        duration_minutes = 60

    title = (data.get("title") or "").strip() or f"{exam.title} — Live Session"
    session_id = str(uuid4())
    pin = _generate_unique_pin()
    now_dt = timezone.now()
    ends_dt = now_dt + timedelta(minutes=duration_minutes)

    with connection.cursor() as cursor:
        cursor.execute(
            """
            INSERT INTO "LiveExamSession"
            ("id", "pin", "exam_id", "group_id", "teacher_id", "title", "simulator_mode", "status", "duration_minutes", "created_at", "ends_at")
            VALUES (%s, %s, %s, %s, %s, %s, %s, 'active', %s, %s, %s)
            """,
            [session_id, pin, exam.id, group_id, actor.id, title, sim_mode, duration_minutes, now_dt, ends_dt],
        )

    return get_session_detail(actor, session_id)


def list_sessions(actor, group_id: str | None = None, status: str | None = None) -> list[dict]:
    ensure_live_exam_tables()
    filters = []
    params = []
    if getattr(actor, "role", None) == "teacher":
        filters.append('"teacher_id"=%s')
        params.append(actor.id)
    if group_id:
        filters.append('"group_id"=%s')
        params.append(group_id)
    if status:
        filters.append('"status"=%s')
        params.append(status)

    where_sql = (" WHERE " + " AND ".join(filters)) if filters else ""
    with connection.cursor() as cursor:
        cursor.execute(
            f"""
            SELECT "id", "pin", "exam_id", "group_id", "teacher_id", "title",
                   "simulator_mode", "status", "duration_minutes", "created_at", "ends_at", "closed_at"
            FROM "LiveExamSession"
            {where_sql}
            ORDER BY "created_at" DESC
            LIMIT 50
            """,
            params,
        )
        rows = cursor.fetchall()
    return [_session_row_to_dict(r) for r in rows]


def get_session_by_pin(pin: str) -> dict:
    ensure_live_exam_tables()
    clean_pin = re.sub(r"\D+", "", str(pin or "").strip())
    if len(clean_pin) != 6:
        raise ContractAPIException("INVALID_PIN", "PIN-kod 6 xonali raqam bo'lishi kerak", 400)

    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT "id", "pin", "exam_id", "group_id", "teacher_id", "title",
                   "simulator_mode", "status", "duration_minutes", "created_at", "ends_at", "closed_at"
            FROM "LiveExamSession"
            WHERE "pin"=%s
            ORDER BY "created_at" DESC
            LIMIT 1
            """,
            [clean_pin],
        )
        row = cursor.fetchone()
    if not row:
        raise ContractAPIException("LIVE_ROOM_NOT_FOUND", "Bunday PIN-kodli imtihon xonasi topilmadi", 404)
    return _session_row_to_dict(row)


def get_session_detail(actor, room_id: str) -> dict:
    ensure_live_exam_tables()
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT "id", "pin", "exam_id", "group_id", "teacher_id", "title",
                   "simulator_mode", "status", "duration_minutes", "created_at", "ends_at", "closed_at"
            FROM "LiveExamSession"
            WHERE "id"=%s OR "pin"=%s
            ORDER BY "created_at" DESC
            LIMIT 1
            """,
            [room_id, room_id],
        )
        row = cursor.fetchone()
    if not row:
        raise ContractAPIException("LIVE_ROOM_NOT_FOUND", "Jonli imtihon xonasi topilmadi", 404)
    return _session_row_to_dict(row)


def _start_or_resume_live_attempt(student_user, exam: MockExam, duration_minutes: int) -> dict:
    """Start or resume a MockAttempt for a student joining a teacher's Live Exam Room."""
    # Ensure StudentProfile exists for this user
    if not StudentProfile.objects.filter(user_id=student_user.id).exists():
        with connection.cursor() as cursor:
            if connection.vendor == "postgresql":
                cursor.execute(
                    """
                    INSERT INTO "StudentProfile" ("userId","availablePrograms","activeProgram","isApproved","currentPoints","linkCode","createdAt","gameQualified")
                    VALUES (%s, ARRAY['IELTS','MULTILEVEL']::"ExamProgram"[], 'IELTS'::"ExamProgram", true, 0, %s, NOW(), false)
                    ON CONFLICT ("userId") DO NOTHING
                    """,
                    [student_user.id, f"BW{random.randint(10000,99999)}"],
                )
            else:
                cursor.execute(
                    """
                    INSERT OR IGNORE INTO "StudentProfile" ("userId","availablePrograms","activeProgram","isApproved","currentPoints","linkCode","createdAt","gameQualified")
                    VALUES (%s, '["IELTS","MULTILEVEL"]', 'IELTS', 1, 0, %s, CURRENT_TIMESTAMP, 0)
                    """,
                    [student_user.id, f"BW{random.randint(10000,99999)}"],
                )

    try:
        return mock_attempts.start(student_user, exam.id, {"mode": "timed"})
    except ContractAPIException:
        # Teacher PIN grants direct access to this room's exam even if unpublished or different track
        existing = (
            MockAttempt.objects.filter(exam_id=exam.id, student_id=student_user.id, status="in_progress")
            .order_by("-started_at")
            .first()
        )
        if existing:
            return mock_attempts.start_response(existing, exam, True, duration_minutes)

        shaped = mock_catalog.shape_exam(exam, False)
        at = timezone.now()
        deadline = at + timedelta(minutes=max(10, duration_minutes))
        first_skill = shaped["sections"][0]["skill"] if shaped.get("sections") else "reading"
        attempt = MockAttempt.objects.create(
            id=str(uuid4()),
            exam_id=exam.id,
            student_id=student_user.id,
            specification_version=exam.specification_version,
            speaking_profile_version=exam.speaking_profile_version,
            status="in_progress",
            started_at=at,
            mode="timed",
            deadline_at=deadline,
            overall_deadline_at=deadline,
            section_deadlines={first_skill: mock_attempts.iso(deadline)},
            flow_mode="single_skill",
            current_skill=first_skill,
            anti_cheat_count=0,
        )
        return mock_attempts.start_response(attempt, exam, False, duration_minutes)


def join_by_pin(student_user, pin: str, simulator_mode_override: str | None = None) -> dict:
    """Allow a student to enter a 6-digit PIN, auto-start/resume the exam attempt, and register in the live monitor."""
    ensure_live_exam_tables()
    if not student_user or not getattr(student_user, "is_authenticated", False):
        raise ContractAPIException("UNAUTHORIZED", "Avval tizimga kiring", 401)

    session = get_session_by_pin(pin)
    if session["status"] == "completed":
        raise ContractAPIException("LIVE_ROOM_CLOSED", "Bu imtihon xonasi allaqachon yakunlangan", 400)

    exam = MockExam.objects.filter(id=session["examId"]).first()
    if not exam:
        raise ContractAPIException("MOCK_EXAM_NOT_FOUND", "Mock imtihon topilmadi", 404)

    start_payload = _start_or_resume_live_attempt(student_user, exam, session["durationMinutes"])
    attempt_id = start_payload["attemptId"]
    signature = start_payload["sessionSignature"]
    total_q = int(start_payload.get("exam", {}).get("questionCount") or 40)
    saved_answers = start_payload.get("savedAnswers") or {}
    answered_cnt = len([v for v in saved_answers.values() if v and str(v).strip()])
    curr_skill = start_payload.get("currentSkill") or "reading"
    sim_mode = (simulator_mode_override or session["simulatorMode"] or "cdi").lower()

    with connection.cursor() as cursor:
        cursor.execute(
            'SELECT "id" FROM "LiveExamParticipant" WHERE "session_id"=%s AND "student_id"=%s',
            [session["id"], student_user.id],
        )
        existing_row = cursor.fetchone()
        if existing_row:
            cursor.execute(
                """
                UPDATE "LiveExamParticipant"
                SET "attempt_id"=%s, "status"='in_progress', "current_skill"=%s,
                    "answered_count"=%s, "total_questions"=%s, "updated_at"=CURRENT_TIMESTAMP
                WHERE "id"=%s
                """,
                [attempt_id, curr_skill, answered_cnt, total_q, existing_row[0]],
            )
        else:
            cursor.execute(
                """
                INSERT INTO "LiveExamParticipant"
                ("id", "session_id", "student_id", "attempt_id", "status", "current_skill",
                 "current_question_number", "answered_count", "flagged_count", "total_questions",
                 "anti_cheat_count", "joined_at", "updated_at")
                VALUES (%s, %s, %s, %s, 'in_progress', %s, 1, %s, 0, %s, 0, CURRENT_TIMESTAMP, CURRENT_TIMESTAMP)
                """,
                [str(uuid4()), session["id"], student_user.id, attempt_id, curr_skill, answered_cnt, total_q],
            )

    return {
        "session": get_session_detail(student_user, session["id"]),
        "roomId": session["id"],
        "pin": session["pin"],
        "examId": exam.id,
        "attemptId": attempt_id,
        "attemptToken": signature,
        "sessionSignature": signature,
        "simulatorMode": sim_mode,
        "redirectUrl": f"/mock/attempt/{attempt_id}?mode={sim_mode}&room={session['id']}&pin={session['pin']}",
        "attempt": start_payload,
    }


def update_progress(user, room_id_or_pin: str, data: dict) -> dict:
    """Update real-time question number, answered count, and flagged count for a live participant."""
    ensure_live_exam_tables()
    session = get_session_detail(user, room_id_or_pin)
    curr_q = data.get("currentQuestionNumber") or data.get("questionNumber") or 1
    answered = data.get("answeredCount")
    flagged = data.get("flaggedCount")
    skill = data.get("currentSkill") or data.get("skill")
    status = data.get("status")

    sets = ['"updated_at"=CURRENT_TIMESTAMP']
    params = []

    if curr_q is not None:
        sets.append('"current_question_number"=%s')
        params.append(max(1, int(curr_q)))
    if answered is not None:
        sets.append('"answered_count"=%s')
        params.append(max(0, int(answered)))
    if flagged is not None:
        sets.append('"flagged_count"=%s')
        params.append(max(0, int(flagged)))
    if skill:
        sets.append('"current_skill"=%s')
        params.append(str(skill))
    if status in ("in_progress", "completed"):
        sets.append('"status"=%s')
        params.append(status)
        if status == "completed":
            sets.append('"finished_at"=CURRENT_TIMESTAMP')

    params.extend([session["id"], user.id])
    with connection.cursor() as cursor:
        cursor.execute(
            f'UPDATE "LiveExamParticipant" SET {", ".join(sets)} WHERE "session_id"=%s AND "student_id"=%s',
            params,
        )
    return {"saved": True, "roomId": session["id"]}


def close_session(actor, room_id: str) -> dict:
    ensure_live_exam_tables()
    session = get_session_detail(actor, room_id)
    if actor.role not in ("admin", "super_admin") and session["teacherId"] != actor.id:
        raise ContractAPIException("FORBIDDEN", "Faqat xona egasi yoki admin xonani yopa oladi", 403)
    with connection.cursor() as cursor:
        cursor.execute(
            """
            UPDATE "LiveExamSession"
            SET "status"='completed', "closed_at"=CURRENT_TIMESTAMP
            WHERE "id"=%s
            """,
            [session["id"]],
        )
    return get_session_detail(actor, session["id"])


def get_monitor_data(actor, room_id: str) -> dict:
    """Synchronize live participant attempts/answers and return full real-time monitoring table."""
    ensure_live_exam_tables()
    session = get_session_detail(actor, room_id)

    total_exam_questions = (
        MockQuestion.objects.filter(group__section__exam_id=session["examId"]).count() or 40
    )

    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT p."id", p."student_id", u."name", u."phone", p."attempt_id", p."status",
                   p."current_skill", p."current_question_number", p."answered_count",
                   p."flagged_count", p."total_questions", p."anti_cheat_count",
                   p."raw_score", p."band_score", p."multilevel_score", p."cefr_level",
                   p."joined_at", p."finished_at", p."updated_at"
            FROM "LiveExamParticipant" p
            JOIN "User" u ON u."id" = p."student_id"
            WHERE p."session_id"=%s
            ORDER BY p."joined_at" ASC
            """,
            [session["id"]],
        )
        part_rows = cursor.fetchall()

    attempt_ids = [r[4] for r in part_rows if r[4]]
    attempts_by_id = {
        a.id: a
        for a in MockAttempt.objects.filter(id__in=attempt_ids).select_related("exam")
    } if attempt_ids else {}

    answers_by_attempt: dict[str, list] = {}
    if attempt_ids:
        for ans in MockAnswer.objects.filter(attempt_id__in=attempt_ids).select_related("question").order_by("updated_at"):
            if (ans.response and ans.response.strip()) or ans.audio_key:
                answers_by_attempt.setdefault(ans.attempt_id, []).append(ans)

    participants_out: list[dict] = []
    seen_student_ids: set[str] = set()

    for r in part_rows:
        (
            p_id,
            s_id,
            s_name,
            s_phone,
            att_id,
            p_status,
            curr_skill,
            curr_q,
            ans_cnt,
            flag_cnt,
            tot_q,
            cheat_cnt,
            raw_sc,
            band_sc,
            ml_sc,
            cefr_lvl,
            joined_at,
            finished_at,
            updated_at,
        ) = r
        seen_student_ids.add(s_id)
        att = attempts_by_id.get(att_id)
        att_answers = answers_by_attempt.get(att_id, [])

        if att is not None:
            live_ans_cnt = len(att_answers)
            if live_ans_cnt > (ans_cnt or 0):
                ans_cnt = live_ans_cnt
            if att_answers:
                last_q = getattr(att_answers[-1], "question", None)
                if last_q and last_q.number:
                    curr_q = max(int(curr_q or 1), int(last_q.number))
            cheat_cnt = max(int(cheat_cnt or 0), int(att.anti_cheat_count or 0))
            curr_skill = att.current_skill or curr_skill
            if att.status == "completed":
                p_status = "completed"
                finished_at = att.finished_at or att.submitted_at or finished_at
                metrics = extract_attempt_dual_metrics(att)
                raw_sc = metrics["rawScore"]
                band_sc = metrics["ieltsBand"]
                ml_sc = metrics["multilevelScore"]
                cefr_lvl = metrics["cefrLevel"]
                with connection.cursor() as cursor:
                    cursor.execute(
                        """
                        UPDATE "LiveExamParticipant"
                        SET "status"='completed', "answered_count"=%s, "anti_cheat_count"=%s,
                            "raw_score"=%s, "band_score"=%s, "multilevel_score"=%s, "cefr_level"=%s,
                            "finished_at"=COALESCE("finished_at", CURRENT_TIMESTAMP)
                        WHERE "id"=%s
                        """,
                        [ans_cnt, cheat_cnt, raw_sc, band_sc, ml_sc, cefr_lvl, p_id],
                    )

        tot_q = int(tot_q or total_exam_questions or 40)
        progress_pct = min(100, round(((ans_cnt or 0) / max(1, tot_q)) * 100))

        participants_out.append(
            {
                "participantId": p_id,
                "studentId": s_id,
                "studentName": s_name,
                "phone": s_phone,
                "attemptId": att_id,
                "status": p_status,
                "currentSkill": curr_skill or "reading",
                "currentQuestionNumber": int(curr_q or 1),
                "answeredCount": int(ans_cnt or 0),
                "flaggedCount": int(flag_cnt or 0),
                "totalQuestions": tot_q,
                "progressPercent": progress_pct,
                "antiCheatCount": int(cheat_cnt or 0),
                "isFlagged": bool((cheat_cnt or 0) > 0),
                "rawScore": round(float(raw_sc), 1) if raw_sc is not None else None,
                "bandScore": round(float(band_sc), 1) if band_sc is not None else None,
                "multilevelScore": round(float(ml_sc), 1) if ml_sc is not None else None,
                "cefrLevel": cefr_lvl,
                "joinedAt": _iso(joined_at),
                "finishedAt": _iso(finished_at),
                "updatedAt": _iso(updated_at),
            }
        )

    # Also include unjoined group roster students if session is bound to a groupId
    if session.get("groupId"):
        roster = list(
            StudentProfile.objects.filter(group_id=session["groupId"])
            .select_related("user")
            .order_by("user__name")
        )
        for sp in roster:
            if sp.user_id in seen_student_ids:
                continue
            participants_out.append(
                {
                    "participantId": None,
                    "studentId": sp.user_id,
                    "studentName": sp.user.name if sp.user else "Student",
                    "phone": sp.user.phone if sp.user else "",
                    "attemptId": None,
                    "status": "not_joined",
                    "currentSkill": None,
                    "currentQuestionNumber": 0,
                    "answeredCount": 0,
                    "flaggedCount": 0,
                    "totalQuestions": total_exam_questions,
                    "progressPercent": 0,
                    "antiCheatCount": 0,
                    "isFlagged": False,
                    "rawScore": None,
                    "bandScore": None,
                    "multilevelScore": None,
                    "cefrLevel": None,
                    "joinedAt": None,
                    "finishedAt": None,
                    "updatedAt": None,
                }
            )

    in_prog_cnt = sum(1 for p in participants_out if p["status"] == "in_progress")
    comp_cnt = sum(1 for p in participants_out if p["status"] == "completed")
    not_joined_cnt = sum(1 for p in participants_out if p["status"] == "not_joined")
    alerts_cnt = sum(p["antiCheatCount"] for p in participants_out)

    return {
        "session": session,
        "summary": {
            "totalStudents": len(participants_out),
            "inProgressCount": in_prog_cnt,
            "completedCount": comp_cnt,
            "notJoinedCount": not_joined_cnt,
            "totalAntiCheatAlerts": alerts_cnt,
        },
        "participants": participants_out,
        "serverTime": timezone.now().isoformat(),
    }


def export_session_excel(actor, room_id: str) -> tuple[bytes, str]:
    """Generate a styled Excel (.xlsx) report of the Live Exam Room results."""
    monitor = get_monitor_data(actor, room_id)
    session = monitor["session"]
    participants = monitor["participants"]

    wb = Workbook()
    ws = wb.active
    ws.title = "Live Exam Results"

    headers = [
        "#",
        "Student Name",
        "Phone",
        "Status",
        "Current Question",
        "Answered / Total",
        "Raw Score",
        "IELTS Band (0-9)",
        "Multi-Level Score (0-75)",
        "CEFR Level",
        "Anti-Cheat Alerts",
        "Started At",
        "Finished At",
    ]
    ws.append(headers)

    header_fill = PatternFill(start_color="065F46", end_color="065F46", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=11)
    thin_side = Side(border_style="thin", color="D1D5DB")
    cell_border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=1, column=col_idx)
        c.fill = header_fill
        c.font = header_font
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = cell_border

    bands_list: list[float] = []
    ml_list: list[float] = []

    for idx, p in enumerate(participants, start=1):
        if p.get("bandScore") is not None:
            bands_list.append(float(p["bandScore"]))
        if p.get("multilevelScore") is not None:
            ml_list.append(float(p["multilevelScore"]))

        row_vals = [
            idx,
            p.get("studentName") or p.get("name") or "-",
            p.get("phone") or "-",
            (p.get("status") or "-").replace("_", " ").title(),
            f"Q{p['currentQuestionNumber']}" if p.get("currentQuestionNumber") else "-",
            f"{p.get('answeredCount', 0)} / {p.get('totalQuestions', 40)}",
            p["rawScore"] if p.get("rawScore") is not None else "-",
            p["bandScore"] if p.get("bandScore") is not None else "-",
            p["multilevelScore"] if p.get("multilevelScore") is not None else "-",
            p.get("cefrLevel") or "-",
            p.get("antiCheatCount", 0),
            (p.get("joinedAt") or "-")[:19].replace("T", " "),
            (p.get("finishedAt") or "-")[:19].replace("T", " "),
        ]
        ws.append(row_vals)
        r_idx = ws.max_row
        for col_idx in range(1, len(headers) + 1):
            ws.cell(row=r_idx, column=col_idx).border = cell_border

    avg_band = round(sum(bands_list) / len(bands_list), 1) if bands_list else "-"
    avg_ml = round(sum(ml_list) / len(ml_list), 1) if ml_list else "-"

    ws.append(
        [
            "",
            "SESSION AVERAGE",
            f"PIN: {session['pin']}",
            f"{len(participants)} Students",
            "-",
            "-",
            "-",
            avg_band,
            avg_ml,
            "-",
            monitor.get("summary", {}).get("totalAntiCheatAlerts", 0),
            "-",
            timezone.now().strftime("%Y-%m-%d %H:%M"),
        ]
    )
    footer_row = ws.max_row
    footer_fill = PatternFill(start_color="ECFDF5", end_color="ECFDF5", fill_type="solid")
    footer_font = Font(bold=True, color="065F46")
    for col_idx in range(1, len(headers) + 1):
        c = ws.cell(row=footer_row, column=col_idx)
        c.fill = footer_fill
        c.font = footer_font
        c.border = cell_border

    for col_idx in range(1, len(headers) + 1):
        col_letter = get_column_letter(col_idx)
        max_len = max(len(str(ws.cell(row=r, column=col_idx).value or "")) for r in range(1, ws.max_row + 1))
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 13), 40)

    buf = BytesIO()
    wb.save(buf)
    safe_title = re.sub(r"[^a-zA-Z0-9_\-]+", "_", session.get("title") or "live_exam").strip("_") or "live_exam"
    filename = f"bestway_{safe_title}_pin_{session['pin']}.xlsx"
    return buf.getvalue(), filename


def record_telemetry(user, data: dict) -> dict:
    """Record heartbeat, focus state, active question, and strikes telemetry from student client."""
    ensure_live_exam_tables()
    if not user or not getattr(user, "is_authenticated", False):
        raise ContractAPIException("UNAUTHORIZED", "Avval tizimga kiring", 401)

    pin = str(data.get("room_pin") or data.get("roomPin") or data.get("pin") or "").strip()
    attempt_id = str(data.get("attempt_id") or data.get("attemptId") or "").strip()
    current_q = data.get("current_question") or data.get("currentQuestion") or data.get("currentQuestionNumber") or 1
    current_section = str(data.get("current_section") or data.get("currentSection") or data.get("currentSkill") or "reading").strip().lower()
    time_remaining = data.get("time_remaining") or data.get("timeRemaining") or data.get("timeRemainingSeconds")
    strikes_count = int(data.get("strikes_count") or data.get("strikesCount") or 0)
    is_focused = bool(data.get("is_focused") if "is_focused" in data else data.get("isFocused", True))

    session_id = None
    if pin:
        clean_pin = re.sub(r"\D+", "", pin)
        if len(clean_pin) == 6:
            with connection.cursor() as cursor:
                cursor.execute(
                    'SELECT "id", "status" FROM "LiveExamSession" WHERE "pin"=%s ORDER BY "created_at" DESC LIMIT 1',
                    [clean_pin],
                )
                s_row = cursor.fetchone()
                if s_row:
                    session_id, session_status = s_row[0], s_row[1]
                    if session_status == "completed":
                        return {
                            "success": True,
                            "action": "end_exam",
                            "status": "completed",
                            "message": "Imtihon xonasi o'qituvchi tomonidan yakunlangan",
                        }

    # Find participant record
    part_row = None
    with connection.cursor() as cursor:
        if session_id:
            cursor.execute(
                """
                SELECT "id", "session_id", "status", "strikes_count", "anti_cheat_count", "attempt_id"
                FROM "LiveExamParticipant"
                WHERE "session_id"=%s AND "student_id"=%s
                """,
                [session_id, user.id],
            )
            part_row = cursor.fetchone()
        if not part_row and attempt_id:
            cursor.execute(
                """
                SELECT "id", "session_id", "status", "strikes_count", "anti_cheat_count", "attempt_id"
                FROM "LiveExamParticipant"
                WHERE "attempt_id"=%s AND "student_id"=%s
                """,
                [attempt_id, user.id],
            )
            part_row = cursor.fetchone()

    if not part_row:
        if attempt_id:
            att = MockAttempt.objects.filter(id=attempt_id, student_id=user.id).first()
            if att:
                cheat_cnt = max(att.anti_cheat_count or 0, strikes_count)
                if strikes_count > (att.anti_cheat_count or 0):
                    att.anti_cheat_count = cheat_cnt
                    if cheat_cnt >= 3:
                        att.status = "completed"
                    att.save(update_fields=["anti_cheat_count", "status"])
                return {
                    "success": True,
                    "action": "disqualified" if cheat_cnt >= 3 else "continue",
                    "status": "disqualified" if cheat_cnt >= 3 else att.status,
                    "strikesCount": cheat_cnt,
                }
        return {"success": True, "action": "continue", "status": "in_progress", "strikesCount": strikes_count}

    p_id, p_session_id, p_status, p_strikes, p_anti_cheat, p_attempt_id = part_row
    effective_strikes = max(int(p_strikes or 0), int(p_anti_cheat or 0), strikes_count)

    # Check if participant is disqualified
    if p_status == "disqualified" or effective_strikes >= 3:
        effective_status = "disqualified"
        effective_strikes = max(effective_strikes, 3)
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE "LiveExamParticipant"
                SET "status"='disqualified', "strikes_count"=%s, "anti_cheat_count"=%s,
                    "last_heartbeat"=CURRENT_TIMESTAMP, "updated_at"=CURRENT_TIMESTAMP
                WHERE "id"=%s
                """,
                [effective_strikes, effective_strikes, p_id],
            )
        target_att = p_attempt_id or attempt_id
        if target_att:
            MockAttempt.objects.filter(id=target_att).update(status="completed", anti_cheat_count=effective_strikes)
        return {
            "success": True,
            "action": "disqualified",
            "status": "disqualified",
            "strikesCount": effective_strikes,
            "message": "3 marta qoidabuzarlik sodir etildi. Imtihondan chetlatildingiz.",
        }

    new_status = "warning" if effective_strikes > 0 else (p_status or "in_progress")
    time_rem_val = int(time_remaining) if time_remaining is not None else None

    with connection.cursor() as cursor:
        cursor.execute(
            """
            UPDATE "LiveExamParticipant"
            SET "current_question_number"=%s,
                "current_skill"=%s,
                "time_remaining_seconds"=COALESCE(%s, "time_remaining_seconds"),
                "strikes_count"=%s,
                "anti_cheat_count"=%s,
                "is_focused"=%s,
                "status"=%s,
                "last_heartbeat"=CURRENT_TIMESTAMP,
                "updated_at"=CURRENT_TIMESTAMP
            WHERE "id"=%s
            """,
            [
                max(1, int(current_q)),
                current_section,
                time_rem_val,
                effective_strikes,
                effective_strikes,
                is_focused,
                new_status,
                p_id,
            ],
        )

    target_att = p_attempt_id or attempt_id
    if target_att and effective_strikes != (p_anti_cheat or 0):
        MockAttempt.objects.filter(id=target_att).update(anti_cheat_count=effective_strikes)

    return {
        "success": True,
        "action": "continue",
        "status": new_status,
        "strikesCount": effective_strikes,
        "isFocused": is_focused,
    }


def get_proctor_status(actor, room_pin_or_id: str) -> dict:
    """Return real-time participant matrix with strikes, focus, online status, and time remaining."""
    ensure_live_exam_tables()
    session = get_session_detail(actor, room_pin_or_id)

    total_exam_questions = (
        MockQuestion.objects.filter(group__section__exam_id=session["examId"]).count() or 40
    )

    now = timezone.now()
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT p."id", p."student_id", u."name", u."phone", p."attempt_id", p."status",
                   p."current_skill", p."current_question_number", p."answered_count",
                   p."flagged_count", p."total_questions", p."anti_cheat_count",
                   p."strikes_count", p."time_remaining_seconds", p."is_focused", p."last_heartbeat",
                   p."joined_at", p."finished_at", p."updated_at"
            FROM "LiveExamParticipant" p
            JOIN "User" u ON u."id" = p."student_id"
            WHERE p."session_id"=%s
            ORDER BY p."joined_at" ASC
            """,
            [session["id"]],
        )
        rows = cursor.fetchall()

    participants: list[dict] = []
    active_cnt = 0
    warning_cnt = 0
    disqualified_cnt = 0
    completed_cnt = 0
    offline_cnt = 0

    for r in rows:
        (
            p_id, s_id, s_name, s_phone, att_id, p_status,
            curr_skill, curr_q, ans_cnt, flag_cnt, tot_q,
            anti_cheat, strikes, time_rem, is_focused, last_hb,
            joined_at, finished_at, updated_at
        ) = r

        effective_strikes = max(int(strikes or 0), int(anti_cheat or 0))

        is_online = False
        if last_hb:
            if hasattr(last_hb, "tzinfo") and last_hb.tzinfo is None:
                last_hb = timezone.make_aware(last_hb)
            diff_seconds = (now - last_hb).total_seconds()
            is_online = diff_seconds <= 30
        elif updated_at:
            if hasattr(updated_at, "tzinfo") and updated_at.tzinfo is None:
                updated_at = timezone.make_aware(updated_at)
            is_online = (now - updated_at).total_seconds() <= 30

        final_status = p_status or "in_progress"
        if effective_strikes >= 3 or p_status == "disqualified":
            final_status = "disqualified"
            disqualified_cnt += 1
        elif p_status == "completed":
            final_status = "completed"
            completed_cnt += 1
        elif not is_online:
            final_status = "offline"
            offline_cnt += 1
        elif effective_strikes > 0 or (is_focused is False):
            final_status = "warning"
            warning_cnt += 1
        else:
            final_status = "in_progress"
            active_cnt += 1

        tot_q = int(tot_q or total_exam_questions or 40)
        progress_pct = min(100, round(((ans_cnt or 0) / max(1, tot_q)) * 100))

        participants.append({
            "participantId": p_id,
            "studentId": s_id,
            "studentName": s_name,
            "name": s_name,
            "phone": s_phone,
            "attemptId": att_id,
            "status": final_status,
            "currentSkill": curr_skill or "reading",
            "currentSection": curr_skill or "reading",
            "currentQuestionNumber": int(curr_q or 1),
            "currentQuestionIndex": int(curr_q or 1),
            "answeredCount": int(ans_cnt or 0),
            "flaggedCount": int(flag_cnt or 0),
            "totalQuestions": tot_q,
            "progressPercent": progress_pct,
            "strikesCount": effective_strikes,
            "antiCheatCount": effective_strikes,
            "timeRemainingSeconds": int(time_rem) if time_rem is not None else None,
            "isFocused": bool(is_focused) if is_focused is not None else True,
            "isOnline": is_online,
            "joinedAt": _iso(joined_at),
            "finishedAt": _iso(finished_at),
            "lastHeartbeat": _iso(last_hb),
            "updatedAt": _iso(updated_at),
        })

    return {
        "room": session,
        "roomId": session["id"],
        "pin": session["pin"],
        "examTitle": session["examTitle"],
        "examType": session["examType"],
        "simulatorMode": session["simulatorMode"],
        "status": session["status"],
        "durationMinutes": session["durationMinutes"],
        "participants": participants,
        "totalCount": len(participants),
        "activeCount": active_cnt,
        "warningCount": warning_cnt,
        "disqualifiedCount": disqualified_cnt,
        "completedCount": completed_cnt,
        "offlineCount": offline_cnt,
    }


def handle_proctor_action(actor, room_pin_or_id: str, data: dict) -> dict:
    """Execute teacher live invigilation actions (start_exam, disqualify_student, warn_student, end_room)."""
    ensure_live_exam_tables()
    session = get_session_detail(actor, room_pin_or_id)
    if actor.role not in ("admin", "super_admin") and session["teacherId"] != actor.id:
        raise ContractAPIException("FORBIDDEN", "Faqat xona egasi yoki admin buyruq bera oladi", 403)

    action = str(data.get("action") or "").strip().lower()
    student_id = str(data.get("student_id") or data.get("studentId") or "").strip()

    if action == "start_exam":
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE "LiveExamSession"
                SET "status"='active', "exam_state"='in_progress'
                WHERE "id"=%s
                """,
                [session["id"]],
            )
            cursor.execute(
                """
                UPDATE "LiveExamParticipant"
                SET "status"='in_progress', "updated_at"=CURRENT_TIMESTAMP
                WHERE "session_id"=%s AND "status" IN ('joined', 'waiting')
                """,
                [session["id"]],
            )
        return {"success": True, "action": "start_exam", "session": get_session_detail(actor, session["id"])}

    elif action == "disqualify_student":
        if not student_id:
            raise ContractAPIException("VALIDATION_ERROR", "student_id ko'rsatilishi shart", 400)
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE "LiveExamParticipant"
                SET "status"='disqualified', "strikes_count"=3, "anti_cheat_count"=3, "updated_at"=CURRENT_TIMESTAMP
                WHERE "session_id"=%s AND "student_id"=%s
                """,
                [session["id"], student_id],
            )
            cursor.execute(
                """
                SELECT "attempt_id" FROM "LiveExamParticipant"
                WHERE "session_id"=%s AND "student_id"=%s
                """,
                [session["id"], student_id],
            )
            row = cursor.fetchone()
            if row and row[0]:
                MockAttempt.objects.filter(id=row[0]).update(status="completed", anti_cheat_count=3)
        return {"success": True, "action": "disqualify_student", "studentId": student_id, "status": "disqualified"}

    elif action == "warn_student":
        if not student_id:
            raise ContractAPIException("VALIDATION_ERROR", "student_id ko'rsatilishi shart", 400)
        with connection.cursor() as cursor:
            cursor.execute(
                """
                UPDATE "LiveExamParticipant"
                SET "strikes_count"=COALESCE("strikes_count", 0) + 1,
                    "anti_cheat_count"=COALESCE("anti_cheat_count", 0) + 1,
                    "status"=CASE WHEN COALESCE("strikes_count", 0) + 1 >= 3 THEN 'disqualified' ELSE 'warning' END,
                    "updated_at"=CURRENT_TIMESTAMP
                WHERE "session_id"=%s AND "student_id"=%s
                """,
                [session["id"], student_id],
            )
            cursor.execute(
                """
                SELECT "strikes_count", "status", "attempt_id" FROM "LiveExamParticipant"
                WHERE "session_id"=%s AND "student_id"=%s
                """,
                [session["id"], student_id],
            )
            row = cursor.fetchone()
            if row:
                strikes_val, stat_val, att_id_val = row
                if stat_val == "disqualified" and att_id_val:
                    MockAttempt.objects.filter(id=att_id_val).update(status="completed", anti_cheat_count=3)
                return {"success": True, "action": "warn_student", "studentId": student_id, "strikesCount": strikes_val, "status": stat_val}
        return {"success": True, "action": "warn_student", "studentId": student_id}

    elif action == "end_room":
        return close_session(actor, session["id"])

    else:
        raise ContractAPIException("INVALID_ACTION", f"Noma'lum buyruq: {action}", 400)
