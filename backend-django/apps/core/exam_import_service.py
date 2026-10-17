"""Smart Exam Question Parser, Excel/CSV Bulk Importer, and Single Audio Timestamp Service."""
from __future__ import annotations

import csv
from io import BytesIO, StringIO
import json
import math
import re
from uuid import uuid4

from django.db import connection, transaction
from django.utils import timezone
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.legacy_schema.models import (
    MockExam,
    MockQuestion,
    MockQuestionGroup,
    MockSection,
)
from common.api.exceptions import ContractAPIException
from .mock_authoring import (
    AUTO_SKILLS,
    STAFF,
    audit_event,
    bump,
    group_payload,
    new_id,
    now,
    question_payload,
)
from .mock_media import media_url, store_upload
from .mock_parse import build_correct_answers
from .mock_rules import CHOICE_TYPES, QUESTION_TYPES, TEXT_TYPES, canonical_decision

_listening_table_ensured = False


def ensure_listening_config_table() -> None:
    global _listening_table_ensured
    if _listening_table_ensured:
        return
    with connection.cursor() as cursor:
        if connection.vendor == "postgresql":
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS "MockListeningConfig" (
                    "exam_id" VARCHAR(255) PRIMARY KEY,
                    "section_id" VARCHAR(255),
                    "audio_key" TEXT,
                    "audio_url" TEXT,
                    "duration_sec" INTEGER,
                    "timestamps" JSONB,
                    "updated_at" TIMESTAMP WITH TIME ZONE DEFAULT NOW()
                );
                """
            )
        else:
            cursor.execute(
                """
                CREATE TABLE IF NOT EXISTS "MockListeningConfig" (
                    "exam_id" VARCHAR(255) PRIMARY KEY,
                    "section_id" VARCHAR(255),
                    "audio_key" TEXT,
                    "audio_url" TEXT,
                    "duration_sec" INTEGER,
                    "timestamps" JSON,
                    "updated_at" TIMESTAMP DEFAULT CURRENT_TIMESTAMP
                );
                """
            )
    _listening_table_ensured = True


PART_HEADER_RE = re.compile(
    r"^\s*(?:PART|PASSAGE|SECTION|READING\s+PASSAGE|LISTENING\s+PART|TASK|BO['‘’]?LIM|QISM)\s*(\d{1,2})\b[:.\-—\s]*(.*)$",
    re.I,
)
ANSWER_SECTION_HEADER_RE = re.compile(
    r"^\s*(?:ANSWERS?|ANSWER\s+KEYS?|CORRECT\s+ANSWERS?|JAVOBLAR|KALITLAR|KEY)\s*[:\-—]?\s*$",
    re.I,
)
INLINE_ANSWER_RE = re.compile(
    r"(?:\(\s*(?:Answer|Ans|Correct|Key|Javob|To['‘’]?g['‘’]?ri\s+javob)\s*[:=\-]\s*([^)]+)\)|"
    r"\[\s*(?:Answer|Ans|Correct|Key|Javob)\s*[:=\-]\s*([^\]]+)\]|"
    r"(?:^|\s)(?:Answer|Ans|Key|Javob)\s*[:=]\s*([A-Za-z0-9_/\s,\-]+)\s*$)",
    re.I,
)
QUESTION_LINE_RE = re.compile(
    r"^\s*(?:(?:Question|Savol|Q|№|#)\s*(\d{1,3})\s*[.):\-]?|\[(\d{1,3})\]|(\d{1,3})\s*[.):\-])\s+(.+)$",
    re.I,
)
OPTION_LINE_RE = re.compile(
    r"^\s*(\*)?\s*\(?([A-Ha-h])\s*[.):\-]\s*(.+?)\s*(\[(?:correct|to['‘’]?g['‘’]?ri|\*)\])?\s*$",
    re.I,
)
INLINE_OPTIONS_RE = re.compile(
    r"(?:^|\s)(\*?)\(?([A-Ha-h])[.)]\s*([^\n\r]+?)(?=(?:\s+\*?\(?[A-Ha-h][.)]\s)|\s*$)"
)
ROMAN_HEADING_LINE_RE = re.compile(
    r"^\s*\(?((?:x|ix|iv|v?i{1,3}|xi{0,3}))[.)]\s+(.+)$",
    re.I,
)
POOL_OPTION_LINE_RE = re.compile(
    r"^\s*(?:List\s+of\s+(?:Headings|Options|People|Places|Ideas|Findings)|Options\s+Pool| Headings)\s*[:\-—]?\s*$",
    re.I,
)
BLANK_TOKEN_RE = re.compile(r"_{2,}|\.{3,}|…|\[gap(?:\s*\d+)?\]|\(\s*\d+\s*\)\s*_{2,}", re.I)
ANSWER_PAIR_RE = re.compile(
    r"(?:^|[,;\s]+)(?:Q|№|#)?(\d{1,3})\s*(?:[.):\-=]|->)\s*([A-Za-z0-9_/'’\-\s]+?)(?=(?:[,;\s]+(?:Q|№|#)?\d{1,3}\s*(?:[.):\-=]|->))|$)",
    re.I,
)


def _detect_type_hint(text_block: str) -> str | None:
    lower = (text_block or "").lower()
    if "not given" in lower or "not_given" in lower:
        if "true" in lower and "false" in lower:
            return "true_false_notgiven"
        if re.search(r"\byes\b", lower) and re.search(r"\bno\b", lower):
            return "yes_no_notgiven"
    if "list of headings" in lower or "matching headings" in lower or "choose the correct heading" in lower:
        return "matching_headings"
    if "match each" in lower or "matching" in lower or "which paragraph contains" in lower:
        return "matching"
    if "complete the summary" in lower or "summary completion" in lower:
        return "summary_completion"
    if "complete the notes" in lower or "note completion" in lower:
        return "note_completion"
    if "complete the sentences" in lower or "sentence completion" in lower:
        return "sentence_completion"
    return None


def _extract_inline_answer(text: str) -> tuple[str, str | None]:
    match = INLINE_ANSWER_RE.search(text)
    if not match:
        return text, None
    ans = (match.group(1) or match.group(2) or match.group(3) or "").strip()
    cleaned = (text[: match.start()] + " " + text[match.end() :]).strip()
    cleaned = re.sub(r"\s+", " ", cleaned)
    return cleaned, ans or None


def _parse_inline_options(line: str) -> tuple[str, list[str], str | None] | None:
    matches = list(INLINE_OPTIONS_RE.finditer(line))
    if len(matches) < 2:
        return None
    letters = [m.group(2).upper() for m in matches]
    if letters[0] != "A":
        return None
    prefix = line[: matches[0].start()].strip()
    options: list[str] = []
    marked_answer: str | None = None
    for m in matches:
        star = bool(m.group(1))
        letter = m.group(2).upper()
        raw_opt = m.group(3).strip()
        if re.search(r"\[(?:correct|to['‘’]?g['‘’]?ri|\*)\]\s*$", raw_opt, re.I):
            star = True
            raw_opt = re.sub(r"\[(?:correct|to['‘’]?g['‘’]?ri|\*)\]\s*$", "", raw_opt, flags=re.I).strip()
        if raw_opt.startswith("*"):
            star = True
            raw_opt = raw_opt.lstrip("*").strip()
        options.append(raw_opt)
        if star and not marked_answer:
            marked_answer = letter
    return prefix, options, marked_answer


def _infer_part_for_question(number: int, exam_type: str = "ielts_academic", skill: str = "reading") -> int:
    if exam_type == "multilevel":
        if skill == "listening":
            # Multilevel listening parts: 1(1-8), 2(9-14), 3(15-18), 4(19-23), 5(24-29), 6(30-35)
            bounds = [(8, 1), (14, 2), (18, 3), (23, 4), (29, 5), (200, 6)]
            return next(p for limit, p in bounds if number <= limit)
        # Multilevel reading parts: 1(1-6), 2(7-14), 3(15-20), 4(21-29), 5(30-35)
        bounds = [(6, 1), (14, 2), (20, 3), (29, 4), (200, 5)]
        return next(p for limit, p in bounds if number <= limit)
    if skill == "listening":
        return min(4, max(1, ((number - 1) // 10) + 1))
    # IELTS Reading: Passage 1 (1-13), Passage 2 (14-26), Passage 3 (27-40)
    if number <= 13:
        return 1
    if number <= 26:
        return 2
    if number <= 40:
        return 3
    return 4


def parse_smart_exam_text(
    text: str,
    default_skill: str = "reading",
    exam_type: str = "ielts_academic",
    provided_answers: dict | None = None,
) -> dict:
    """Smart parser capable of extracting 1..40+ questions, passages, options, and answer keys."""
    if not isinstance(text, str) or not text.strip():
        raise ContractAPIException("VALIDATION_ERROR", "Matn bo'sh bo'lmasligi kerak", 400)

    lines = re.split(r"\r\n?|\n", text)
    answer_map: dict[int, str] = {}
    if isinstance(provided_answers, dict):
        for k, v in provided_answers.items():
            try:
                q_num = int(str(k).strip())
                if v is not None and str(v).strip():
                    answer_map[q_num] = str(v).strip()
            except ValueError:
                continue

    global_hint = _detect_type_hint(text)
    in_answer_section = False
    current_part = 1
    explicit_parts_seen = False

    parts_meta: dict[int, dict] = {
        1: {
            "partNumber": 1,
            "title": "Passage 1" if default_skill == "reading" else "Part 1",
            "preambleLines": [],
            "sharedOptions": [],
            "hint": None,
        }
    }
    raw_questions: list[dict] = []
    current_q: dict | None = None

    def flush_question():
        nonlocal current_q
        if current_q is not None:
            raw_questions.append(current_q)
            current_q = None

    for raw_line in lines:
        line = raw_line.strip()
        if not line:
            continue

        if ANSWER_SECTION_HEADER_RE.match(line):
            flush_question()
            in_answer_section = True
            continue

        if in_answer_section:
            pairs = list(ANSWER_PAIR_RE.finditer(line))
            if pairs:
                for m in pairs:
                    q_num = int(m.group(1))
                    ans_val = m.group(2).strip()
                    if ans_val:
                        answer_map[q_num] = ans_val
                continue
            single_m = re.match(r"^\s*(?:Q|№|#)?(\d{1,3})\s*[.):\-=]+\s*(.+)$", line, re.I)
            if single_m:
                answer_map[int(single_m.group(1))] = single_m.group(2).strip()
                continue

        part_match = PART_HEADER_RE.match(line)
        if part_match:
            flush_question()
            in_answer_section = False
            explicit_parts_seen = True
            current_part = int(part_match.group(1))
            rest_title = (part_match.group(2) or "").strip()
            unit = "Passage" if default_skill == "reading" and exam_type != "multilevel" else "Part"
            parts_meta.setdefault(
                current_part,
                {
                    "partNumber": current_part,
                    "title": f"{unit} {current_part}" + (f" — {rest_title}" if rest_title else ""),
                    "preambleLines": [],
                    "sharedOptions": [],
                    "hint": None,
                },
            )
            continue

        # Check inline options on a single line (either inside current question or as part of question line)
        if current_q is not None:
            inline_opts = _parse_inline_options(line)
            if inline_opts and not inline_opts[0]:
                _, opts, marked = inline_opts
                current_q["options"].extend(opts)
                if marked and not current_q.get("inlineAnswer"):
                    current_q["inlineAnswer"] = marked
                continue

            opt_match = OPTION_LINE_RE.match(line)
            if opt_match:
                star = bool(opt_match.group(1) or opt_match.group(4))
                letter = opt_match.group(2).upper()
                opt_text = opt_match.group(3).strip()
                opt_text, inline_ans = _extract_inline_answer(opt_text)
                if inline_ans and not current_q.get("inlineAnswer"):
                    current_q["inlineAnswer"] = inline_ans
                if star and not current_q.get("inlineAnswer"):
                    current_q["inlineAnswer"] = letter
                current_q["options"].append(opt_text)
                continue

        # Check if line starts a new question
        q_match = QUESTION_LINE_RE.match(line)
        if q_match:
            q_num = int(q_match.group(1) or q_match.group(2) or q_match.group(3))
            q_body = q_match.group(4).strip()
            # Ensure this isn't just an answer key line at the bottom like "1. A" when we already have question 1
            existing_nums = {rq["number"] for rq in raw_questions} | ({current_q["number"]} if current_q else set())
            if q_num in existing_nums and len(q_body) <= 18 and re.fullmatch(
                r"[A-Ha-h]|TRUE|FALSE|NOT\s+GIVEN|YES|NO|T|F|NG|Y|N|[ivxIVX]{1,4}", q_body, re.I
            ):
                flush_question()
                answer_map[q_num] = q_body.strip()
                continue

            flush_question()
            q_body, inline_ans = _extract_inline_answer(q_body)
            inline_opts = _parse_inline_options(q_body)
            opts_list: list[str] = []
            if inline_opts and inline_opts[0]:
                q_body, opts_list, marked = inline_opts
                if marked and not inline_ans:
                    inline_ans = marked

            target_part = (
                current_part
                if explicit_parts_seen
                else _infer_part_for_question(q_num, exam_type=exam_type, skill=default_skill)
            )
            parts_meta.setdefault(
                target_part,
                {
                    "partNumber": target_part,
                    "title": ("Passage " if default_skill == "reading" and exam_type != "multilevel" else "Part ")
                    + str(target_part),
                    "preambleLines": [],
                    "sharedOptions": [],
                    "hint": None,
                },
            )
            current_q = {
                "number": q_num,
                "partNumber": target_part,
                "lines": [q_body],
                "options": opts_list,
                "inlineAnswer": inline_ans,
            }
            continue

        # Check Roman numeral heading lines or option pool lines in preamble before questions
        if current_q is None:
            p_meta = parts_meta.setdefault(
                current_part,
                {
                    "partNumber": current_part,
                    "title": f"Part {current_part}",
                    "preambleLines": [],
                    "sharedOptions": [],
                    "hint": None,
                },
            )
            roman_m = ROMAN_HEADING_LINE_RE.match(line)
            opt_pool_m = OPTION_LINE_RE.match(line)
            if roman_m:
                p_meta["sharedOptions"].append(f"{roman_m.group(1).lower()}. {roman_m.group(2).strip()}")
                p_meta["hint"] = "matching_headings"
            elif opt_pool_m and not POOL_OPTION_LINE_RE.match(line):
                p_meta["sharedOptions"].append(opt_pool_m.group(3).strip())
                if not p_meta["hint"]:
                    p_meta["hint"] = "matching"
            else:
                p_meta["preambleLines"].append(line)
                detected = _detect_type_hint(line)
                if detected:
                    p_meta["hint"] = detected
        else:
            cleaned_line, inline_ans = _extract_inline_answer(line)
            if inline_ans and not current_q.get("inlineAnswer"):
                current_q["inlineAnswer"] = inline_ans
            if cleaned_line:
                current_q["lines"].append(cleaned_line)

    flush_question()

    # Build normalized questions and parts
    questions_out: list[dict] = []
    by_type_summary: dict[str, int] = {}
    with_answers_count = 0

    for rq in raw_questions:
        q_num = rq["number"]
        part_num = rq["partNumber"]
        p_meta = parts_meta.get(part_num) or {}
        part_hint = p_meta.get("hint") or global_hint
        shared_options = p_meta.get("sharedOptions") or []

        prompt = re.sub(r"\s+", " ", " ".join(rq["lines"])).strip()
        raw_ans = rq.get("inlineAnswer") or answer_map.get(q_num)
        if raw_ans:
            answer_map[q_num] = str(raw_ans).strip()

        options = list(rq["options"]) if rq["options"] else []
        ans_upper = (str(raw_ans).strip().upper() if raw_ans else "")

        # Determine question type
        if len(options) >= 2:
            q_type = "multiple_choice"
        elif ans_upper in {"TRUE", "FALSE", "NOT GIVEN", "NOT_GIVEN", "T", "F", "NG"} or part_hint == "true_false_notgiven":
            q_type = "true_false_notgiven"
            options = ["TRUE", "FALSE", "NOT GIVEN"]
        elif ans_upper in {"YES", "NO", "Y", "N"} or part_hint == "yes_no_notgiven":
            q_type = "yes_no_notgiven"
            options = ["YES", "NO", "NOT GIVEN"]
        elif part_hint in {"matching_headings", "matching"} and shared_options:
            q_type = part_hint
            options = list(shared_options)
        elif re.fullmatch(r"(?:x|ix|iv|v?i{1,3}|xi{0,3})", ans_upper, re.I):
            q_type = "matching_headings"
            options = list(shared_options) if len(shared_options) >= 2 else [
                "i", "ii", "iii", "iv", "v", "vi", "vii", "viii", "ix", "x"
            ]
        elif re.fullmatch(r"[A-Ha-h]", ans_upper) and not options:
            if len(shared_options) >= 2:
                q_type = part_hint if part_hint in {"matching", "matching_headings"} else "matching"
                options = list(shared_options)
            else:
                q_type = "multiple_choice"
                options = ["Option A", "Option B", "Option C", "Option D"]
                idx = ord(ans_upper) - ord("A")
                while len(options) <= idx:
                    options.append(f"Option {chr(65 + len(options))}")
        elif part_hint in {"summary_completion", "note_completion", "sentence_completion"}:
            q_type = part_hint
            options = None
        elif BLANK_TOKEN_RE.search(prompt):
            q_type = "sentence_completion"
            options = None
        else:
            q_type = "short_answer"
            options = None

        correct_answers: list[str] = []
        accepted_variants: list[str] = []
        if raw_ans:
            built = build_correct_answers(q_type, options, str(raw_ans))
            if q_type in {"true_false_notgiven", "yes_no_notgiven"}:
                canon = [
                    canonical_decision(x)
                    for x in built
                    if canonical_decision(x) in {"TRUE", "FALSE", "YES", "NO", "NOT_GIVEN"}
                ]
                correct_answers = [ canon[0] ] if canon else [canonical_decision(str(raw_ans))]
            elif q_type in CHOICE_TYPES:
                correct_answers = built if built else [str(raw_ans).strip()]
            else:
                if built:
                    correct_answers = [built[0]]
                    accepted_variants = built[1:]
                else:
                    correct_answers = [str(raw_ans).strip()]

        if correct_answers:
            with_answers_count += 1

        word_limit = None
        answer_rule = None
        if q_type in TEXT_TYPES:
            if exam_type == "multilevel":
                word_limit = 1
                answer_rule = "ONE_WORD_AND_OR_NUMBER"
            else:
                word_limit = 3

        q_item = {
            "number": q_num,
            "partNumber": part_num,
            "type": q_type,
            "prompt": prompt,
            "options": options if options else None,
            "correctAnswers": correct_answers,
            "acceptedVariants": accepted_variants,
            "points": 1,
            "wordLimit": word_limit,
            "answerRule": answer_rule,
        }
        questions_out.append(q_item)
        by_type_summary[q_type] = by_type_summary.get(q_type, 0) + 1

    # Assemble parts list
    parts_out: list[dict] = []
    for p_num in sorted(parts_meta.keys()):
        pm = parts_meta[p_num]
        p_qs = [q for q in questions_out if q["partNumber"] == p_num]
        if not p_qs and not pm["preambleLines"]:
            continue
        preamble_text = "\n".join(pm["preambleLines"]).strip() or None
        instructions_text = None
        passage_text = None
        if preamble_text:
            if len(preamble_text) > 220 and default_skill == "reading":
                passage_text = preamble_text
            else:
                instructions_text = preamble_text
        parts_out.append(
            {
                "partNumber": p_num,
                "title": pm["title"],
                "passageText": passage_text,
                "instructions": instructions_text,
                "questions": p_qs,
            }
        )

    first_instructions = next((p["instructions"] for p in parts_out if p.get("instructions")), None)
    first_passage = next((p["passageText"] for p in parts_out if p.get("passageText")), None)

    return {
        "instructions": first_instructions,
        "passageText": first_passage,
        "parts": parts_out,
        "questions": questions_out,
        "answers": {str(k): v for k, v in answer_map.items()},
        "count": len(questions_out),
        "summary": {
            "total": len(questions_out),
            "withAnswers": with_answers_count,
            "byType": by_type_summary,
        },
    }


def _normalize_incoming_question(raw_q: dict, default_num: int, exam_type: str, skill: str) -> dict:
    number = int(raw_q.get("number") or raw_q.get("questionNumber") or default_num)
    q_type = str(raw_q.get("type") or raw_q.get("questionType") or "short_answer").strip()
    if q_type not in QUESTION_TYPES:
        q_type = "short_answer"
    prompt = str(raw_q.get("prompt") or "").strip() or f"Question {number}"
    options = raw_q.get("options")
    if isinstance(options, list):
        options = [str(o).strip() for o in options if str(o).strip()]
    else:
        options = None

    if q_type == "true_false_notgiven" and not options:
        options = ["TRUE", "FALSE", "NOT GIVEN"]
    elif q_type == "yes_no_notgiven" and not options:
        options = ["YES", "NO", "NOT GIVEN"]
    elif q_type in TEXT_TYPES:
        options = None

    correct_answers = raw_q.get("correctAnswers")
    if isinstance(correct_answers, str) and correct_answers.strip():
        correct_answers = build_correct_answers(q_type, options, correct_answers.strip())
    elif isinstance(correct_answers, list):
        correct_answers = [str(x).strip() for x in correct_answers if str(x).strip()]
    else:
        single_ans = raw_q.get("correctAnswer") or raw_q.get("answer")
        if single_ans is not None and str(single_ans).strip():
            correct_answers = build_correct_answers(q_type, options, str(single_ans).strip())
        else:
            correct_answers = []

    if q_type in {"true_false_notgiven", "yes_no_notgiven"} and correct_answers:
        correct_answers = [canonical_decision(correct_answers[0])]

    accepted_variants = raw_q.get("acceptedVariants")
    if isinstance(accepted_variants, str):
        accepted_variants = [x.strip() for x in re.split(r"[/;,]+", accepted_variants) if x.strip()]
    elif isinstance(accepted_variants, list):
        accepted_variants = [str(x).strip() for x in accepted_variants if str(x).strip()]
    else:
        accepted_variants = []

    if q_type not in TEXT_TYPES:
        accepted_variants = []

    points = int(raw_q.get("points") or 1)
    word_limit = raw_q.get("wordLimit")
    if word_limit is not None and str(word_limit).strip() != "":
        try:
            word_limit = int(word_limit)
        except ValueError:
            word_limit = None
    elif q_type in TEXT_TYPES and exam_type == "multilevel":
        word_limit = 1

    answer_rule = raw_q.get("answerRule")
    if not answer_rule and q_type in TEXT_TYPES and exam_type == "multilevel":
        answer_rule = "ONE_WORD_AND_OR_NUMBER"

    return {
        "number": number,
        "partNumber": int(
            raw_q.get("partNumber")
            or raw_q.get("part")
            or _infer_part_for_question(number, exam_type=exam_type, skill=skill)
        ),
        "type": q_type,
        "prompt": prompt,
        "options": options,
        "correct_answers": correct_answers,
        "accepted_variants": accepted_variants,
        "points": max(1, points),
        "word_limit": word_limit,
        "answer_rule": answer_rule,
    }


def _assert_staff_author(actor, exam: MockExam) -> None:
    if not actor or getattr(actor, "role", None) not in STAFF:
        raise ContractAPIException("FORBIDDEN", "Bu amal uchun rolingiz yetarli emas", 403)
    if actor.role == "teacher" and exam.created_by_id and exam.created_by_id != actor.id:
        raise ContractAPIException(
            "MOCK_NOT_OWNER",
            "Bu imtihonni faqat yaratgan o'qituvchi (yoki admin) tahrirlay oladi",
            403,
        )


def bulk_import_to_group(actor, group_id: str, payload_data: dict) -> dict:
    """Persist smart-parsed questions or structured question array into a MockQuestionGroup."""
    try:
        group = MockQuestionGroup.objects.select_related("section__exam").get(id=group_id)
    except MockQuestionGroup.DoesNotExist:
        raise ContractAPIException("MOCK_GROUP_NOT_FOUND", "Savollar guruhi topilmadi", 404)

    exam = group.section.exam
    skill = group.section.skill
    _assert_staff_author(actor, exam)

    replace_existing = bool(payload_data.get("replaceExisting", False))
    incoming_questions = payload_data.get("questions")
    parsed_meta = {}

    if not isinstance(incoming_questions, list):
        raw_text = payload_data.get("text") or ""
        parsed_meta = parse_smart_exam_text(
            raw_text,
            default_skill=skill,
            exam_type=exam.type,
            provided_answers=payload_data.get("answers"),
        )
        incoming_questions = parsed_meta["questions"]

    if not incoming_questions:
        raise ContractAPIException("VALIDATION_ERROR", "Hech qanday savol topilmadi", 400)

    with transaction.atomic():
        locked = MockExam.objects.select_for_update().get(id=exam.id)
        group = MockQuestionGroup.objects.select_for_update().get(id=group_id)

        passage_text = payload_data.get("passageText") or parsed_meta.get("passageText")
        instructions = payload_data.get("instructions") or parsed_meta.get("instructions")
        title = payload_data.get("title")
        update_group_fields = []
        if isinstance(passage_text, str) and passage_text.strip():
            group.passage_text = passage_text.strip()
            update_group_fields.append("passage_text")
        if isinstance(instructions, str) and instructions.strip():
            group.instructions = instructions.strip()
            update_group_fields.append("instructions")
        if isinstance(title, str) and title.strip():
            group.title = title.strip()
            update_group_fields.append("title")
        if update_group_fields:
            group.save(update_fields=update_group_fields)

        if replace_existing:
            MockQuestion.objects.filter(group_id=group.id).delete()

        existing_by_num = {
            q.number: q for q in MockQuestion.objects.filter(group_id=group.id).order_by("sort_order", "number")
        }
        next_sort = len(existing_by_num)
        saved_questions: list[MockQuestion] = []

        for idx, raw_q in enumerate(incoming_questions):
            norm = _normalize_incoming_question(raw_q, idx + 1, locked.type, skill)
            q_num = norm["number"]
            if q_num in existing_by_num:
                q_obj = existing_by_num[q_num]
                for field in (
                    "type",
                    "prompt",
                    "options",
                    "correct_answers",
                    "accepted_variants",
                    "points",
                    "word_limit",
                    "answer_rule",
                ):
                    setattr(q_obj, field, norm[field])
                q_obj.save(
                    update_fields=[
                        "type",
                        "prompt",
                        "options",
                        "correct_answers",
                        "accepted_variants",
                        "points",
                        "word_limit",
                        "answer_rule",
                    ]
                )
            else:
                q_obj = MockQuestion.objects.create(
                    id=new_id(),
                    group_id=group.id,
                    number=q_num,
                    sort_order=next_sort,
                    type=norm["type"],
                    prompt=norm["prompt"],
                    options=norm["options"],
                    correct_answers=norm["correct_answers"],
                    accepted_variants=norm["accepted_variants"],
                    points=norm["points"],
                    word_limit=norm["word_limit"],
                    answer_rule=norm["answer_rule"],
                    created_at=now(),
                )
                next_sort += 1
            saved_questions.append(q_obj)

        bump(locked)

    audit_event(
        actor,
        "mock.group.bulk_import",
        "mockQuestionGroup",
        group.id,
        new={"imported": len(saved_questions), "replaceExisting": replace_existing},
    )
    all_group_questions = list(MockQuestion.objects.filter(group_id=group.id).order_by("sort_order", "number"))
    return {
        "imported": len(saved_questions),
        "groupId": group.id,
        "examId": exam.id,
        "version": locked.content_version,
        "group": group_payload(group, questions=all_group_questions),
        "questions": [question_payload(q) for q in saved_questions],
    }


def bulk_import_to_section(actor, section_id: str, payload_data: dict) -> dict:
    try:
        section = MockSection.objects.select_related("exam").get(id=section_id)
    except MockSection.DoesNotExist:
        raise ContractAPIException("MOCK_SECTION_NOT_FOUND", "Bo'lim topilmadi", 404)
    return bulk_import_to_exam(actor, section.exam_id, section.skill, payload_data)


def bulk_import_to_exam(actor, exam_id: str, skill: str, payload_data: dict) -> dict:
    """Bulk import up to 40+ questions into an exam's skill section, auto-distributing across parts/passages."""
    try:
        exam = MockExam.objects.get(id=exam_id)
    except MockExam.DoesNotExist:
        raise ContractAPIException("MOCK_EXAM_NOT_FOUND", "Mock imtihon topilmadi", 404)

    _assert_staff_author(actor, exam)
    skill = (skill or payload_data.get("skill") or "reading").strip().lower()
    if skill not in {"listening", "reading", "writing", "speaking"}:
        raise ContractAPIException("VALIDATION_ERROR", "Noto'g'ri skill nomi", 400)

    group_id = payload_data.get("groupId")
    if group_id:
        return bulk_import_to_group(actor, group_id, payload_data)

    replace_existing = bool(payload_data.get("replaceExisting", True))
    incoming_questions = payload_data.get("questions")
    parts_meta_by_num: dict[int, dict] = {}

    if not isinstance(incoming_questions, list):
        raw_text = payload_data.get("text") or ""
        parsed = parse_smart_exam_text(
            raw_text,
            default_skill=skill,
            exam_type=exam.type,
            provided_answers=payload_data.get("answers"),
        )
        incoming_questions = parsed["questions"]
        for p in parsed.get("parts") or []:
            parts_meta_by_num[p["partNumber"]] = p
    else:
        for p in payload_data.get("parts") or []:
            if isinstance(p, dict) and p.get("partNumber"):
                parts_meta_by_num[int(p["partNumber"])] = p

    if not incoming_questions:
        raise ContractAPIException("VALIDATION_ERROR", "Import qilish uchun savollar topilmadi", 400)

    with transaction.atomic():
        locked = MockExam.objects.select_for_update().get(id=exam.id)
        section = MockSection.objects.filter(exam_id=locked.id, skill=skill).first()
        if section is None:
            sort_idx = {"listening": 0, "reading": 1, "writing": 2, "speaking": 3}.get(skill, 0)
            section = MockSection.objects.create(
                id=new_id(),
                exam_id=locked.id,
                skill=skill,
                title=skill.title(),
                sort_order=sort_idx,
                duration_minutes=40 if skill == "listening" else 60,
                instructions=payload_data.get("instructions"),
            )

        existing_groups = list(MockQuestionGroup.objects.filter(section_id=section.id).order_by("sort_order", "id"))
        normalized_list = [
            _normalize_incoming_question(rq, idx + 1, locked.type, skill)
            for idx, rq in enumerate(incoming_questions)
        ]

        part_numbers = sorted({nq["partNumber"] for nq in normalized_list})
        groups_by_part: dict[int, MockQuestionGroup] = {}
        unit_name = "Passage" if skill == "reading" and locked.type != "multilevel" else "Part"

        for idx, p_num in enumerate(part_numbers):
            matched_group = next(
                (g for g in existing_groups if g.part_number == p_num or g.sort_order == p_num - 1),
                None,
            )
            if matched_group is None and idx < len(existing_groups):
                matched_group = existing_groups[idx]
            if matched_group is None:
                p_info = parts_meta_by_num.get(p_num) or {}
                matched_group = MockQuestionGroup.objects.create(
                    id=new_id(),
                    section_id=section.id,
                    sort_order=max(0, p_num - 1),
                    title=p_info.get("title") or f"{unit_name} {p_num}",
                    instructions=p_info.get("instructions") or payload_data.get("instructions"),
                    passage_text=p_info.get("passageText") or payload_data.get("passageText"),
                    part_number=p_num if skill == "listening" else None,
                    audio_play_limit=2 if locked.type == "multilevel" and skill == "listening" else 1,
                    created_at=now(),
                )
            else:
                p_info = parts_meta_by_num.get(p_num) or {}
                upd = []
                if p_info.get("passageText") and not matched_group.passage_text:
                    matched_group.passage_text = p_info["passageText"]
                    upd.append("passage_text")
                if p_info.get("instructions") and not matched_group.instructions:
                    matched_group.instructions = p_info["instructions"]
                    upd.append("instructions")
                if skill == "listening" and matched_group.part_number != p_num:
                    matched_group.part_number = p_num
                    upd.append("part_number")
                if upd:
                    matched_group.save(update_fields=upd)
            groups_by_part[p_num] = matched_group

        if replace_existing:
            target_group_ids = [g.id for g in groups_by_part.values()]
            MockQuestion.objects.filter(group_id__in=target_group_ids).delete()

        saved_questions: list[MockQuestion] = []
        sort_counters: dict[str, int] = {}

        for norm in normalized_list:
            grp = groups_by_part[norm["partNumber"]]
            s_idx = sort_counters.get(grp.id, 0)
            sort_counters[grp.id] = s_idx + 1

            existing_q = MockQuestion.objects.filter(group_id=grp.id, number=norm["number"]).first()
            if existing_q:
                for field in (
                    "type",
                    "prompt",
                    "options",
                    "correct_answers",
                    "accepted_variants",
                    "points",
                    "word_limit",
                    "answer_rule",
                ):
                    setattr(existing_q, field, norm[field])
                existing_q.save(
                    update_fields=[
                        "type",
                        "prompt",
                        "options",
                        "correct_answers",
                        "accepted_variants",
                        "points",
                        "word_limit",
                        "answer_rule",
                    ]
                )
                saved_questions.append(existing_q)
            else:
                q_obj = MockQuestion.objects.create(
                    id=new_id(),
                    group_id=grp.id,
                    number=norm["number"],
                    sort_order=s_idx,
                    type=norm["type"],
                    prompt=norm["prompt"],
                    options=norm["options"],
                    correct_answers=norm["correct_answers"],
                    accepted_variants=norm["accepted_variants"],
                    points=norm["points"],
                    word_limit=norm["word_limit"],
                    answer_rule=norm["answer_rule"],
                    created_at=now(),
                )
                saved_questions.append(q_obj)

        bump(locked)

    audit_event(
        actor,
        "mock.exam.bulk_import",
        "mockExam",
        exam.id,
        new={"skill": skill, "imported": len(saved_questions), "parts": len(groups_by_part)},
    )
    return {
        "imported": len(saved_questions),
        "examId": exam.id,
        "sectionId": section.id,
        "skill": skill,
        "partsUpdated": len(groups_by_part),
        "version": locked.content_version,
        "questions": [question_payload(q) for q in saved_questions],
    }


EXCEL_HEADERS = [
    "Part",
    "QuestionNumber",
    "QuestionType",
    "Prompt",
    "OptionA",
    "OptionB",
    "OptionC",
    "OptionD",
    "OptionE",
    "OptionF",
    "OptionG",
    "OptionH",
    "CorrectAnswer",
    "AcceptedVariants",
    "Points",
    "WordLimit",
    "PassageText",
    "Instructions",
    "AudioStartSec",
    "AudioEndSec",
]


def generate_excel_template(exam_type: str = "ielts_academic", skill: str = "reading") -> tuple[bytes, str]:
    """Generate a styled 40-question .xlsx template for teachers to fill and upload."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Questions Import"
    ws.append(EXCEL_HEADERS)

    header_fill = PatternFill(start_color="0F172A", end_color="0F172A", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF", size=10)
    thin_side = Side(border_style="thin", color="CBD5E1")
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)

    for col_idx in range(1, len(EXCEL_HEADERS) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = border

    total_rows = 35 if exam_type == "multilevel" else 40
    for q_num in range(1, total_rows + 1):
        part_num = _infer_part_for_question(q_num, exam_type=exam_type, skill=skill)
        is_first_in_part = q_num == 1 or _infer_part_for_question(q_num - 1, exam_type=exam_type, skill=skill) != part_num
        if q_num <= 10:
            q_type = "multiple_choice"
            prompt = f"Sample question {q_num}: Choose the correct letter A, B, C or D."
            opts = ["First option", "Second option", "Third option", "Fourth option", "", "", "", ""]
            ans = "A"
            variants = ""
            w_limit = ""
        elif q_num <= 20:
            q_type = "true_false_notgiven" if skill == "reading" else "short_answer"
            if q_type == "true_false_notgiven":
                prompt = f"Sample statement {q_num}: The research was conducted in 2025."
                opts = ["TRUE", "FALSE", "NOT GIVEN", "", "", "", "", ""]
                ans = "TRUE"
                variants = ""
                w_limit = ""
            else:
                prompt = f"Sample note completion {q_num}: The train leaves at ____."
                opts = ["", "", "", "", "", "", "", ""]
                ans = "9:30"
                variants = "nine thirty"
                w_limit = 2
        elif q_num <= 30:
            q_type = "matching"
            prompt = f"Sample matching item {q_num}: Match the statement with the correct option."
            opts = ["Option A", "Option B", "Option C", "Option D", "Option E", "Option F", "", ""]
            ans = "B"
            variants = ""
            w_limit = ""
        else:
            q_type = "sentence_completion"
            prompt = f"Sample completion {q_num}: Researchers discovered a new ____ in the region."
            opts = ["", "", "", "", "", "", "", ""]
            ans = "species"
            variants = "plant species"
            w_limit = 1 if exam_type == "multilevel" else 2

        passage = (
            f"Enter full reading passage text for Part/Passage {part_num} here (only needed on the first row of each part)."
            if is_first_in_part and skill == "reading"
            else ""
        )
        instr = (
            f"Questions for Part {part_num}."
            if is_first_in_part
            else ""
        )
        start_sec = (part_num - 1) * 420 if is_first_in_part and skill == "listening" else ""
        end_sec = part_num * 420 if is_first_in_part and skill == "listening" else ""

        row_values = [
            part_num,
            q_num,
            q_type,
            prompt,
            *opts,
            ans,
            variants,
            1,
            w_limit,
            passage,
            instr,
            start_sec,
            end_sec,
        ]
        ws.append(row_values)

    for col_idx in range(1, len(EXCEL_HEADERS) + 1):
        col_letter = get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = 28 if col_idx in (4, 17, 18) else 15

    buf = BytesIO()
    wb.save(buf)
    filename = f"bestway_{exam_type}_{skill}_template.xlsx"
    return buf.getvalue(), filename


def parse_spreadsheet_bytes(file_bytes: bytes, filename: str = "import.xlsx") -> dict:
    """Parse an uploaded .xlsx or .csv file into parts, questions, and listening timestamps."""
    if not file_bytes:
        raise ContractAPIException("VALIDATION_ERROR", "Yuklangan fayl bo'sh", 400)

    lower_name = (filename or "import.xlsx").lower()
    rows_dicts: list[dict[str, str]] = []

    if lower_name.endswith(".csv"):
        try:
            decoded = file_bytes.decode("utf-8-sig")
        except UnicodeDecodeError:
            decoded = file_bytes.decode("latin-1")
        reader = csv.DictReader(StringIO(decoded))
        for r in reader:
            rows_dicts.append(
                {
                    str(k).strip(): (str(v[0]) if isinstance(v, list) and v else str(v or "")).strip()
                    for k, v in r.items()
                    if k is not None and str(k).strip()
                }
            )
    else:
        try:
            wb = load_workbook(BytesIO(file_bytes), data_only=True)
        except Exception as exc:
            raise ContractAPIException("VALIDATION_ERROR", f"Excel faylni o'qib bo'lmadi: {exc}", 400)
        ws = wb.active
        all_rows = list(ws.iter_rows(values_only=True))
        if len(all_rows) < 2:
            raise ContractAPIException("VALIDATION_ERROR", "Jadvalda savollar qatori topilmadi", 400)
        headers = [str(h or "").strip() for h in all_rows[0]]
        for raw_row in all_rows[1:]:
            if not any(cell is not None and str(cell).strip() != "" for cell in raw_row):
                continue
            row_map = {}
            for idx, h in enumerate(headers):
                if h and idx < len(raw_row) and raw_row[idx] is not None:
                    row_map[h] = str(raw_row[idx]).strip()
            if row_map:
                rows_dicts.append(row_map)

    if not rows_dicts:
        raise ContractAPIException("VALIDATION_ERROR", "Jadvalda hech qanday savol topilmadi", 400)

    def _get(row: dict, *keys: str) -> str:
        lower_lookup = {k.lower().replace(" ", "").replace("_", ""): v for k, v in row.items()}
        for k in keys:
            norm_k = k.lower().replace(" ", "").replace("_", "")
            if norm_k in lower_lookup and lower_lookup[norm_k] != "":
                return lower_lookup[norm_k]
        return ""

    questions_out: list[dict] = []
    parts_map: dict[int, dict] = {}
    timestamps_map: dict[int, dict] = {}
    by_type: dict[str, int] = {}
    with_answers = 0

    for idx, r in enumerate(rows_dicts):
        q_num_raw = _get(r, "QuestionNumber", "Number", "Q", "No", "SavolRaqami")
        try:
            q_num = int(float(q_num_raw)) if q_num_raw else (idx + 1)
        except ValueError:
            q_num = idx + 1

        part_raw = _get(r, "Part", "PartNumber", "Passage", "Section", "Qism")
        try:
            part_num = int(float(part_raw)) if part_raw else _infer_part_for_question(q_num)
        except ValueError:
            part_num = _infer_part_for_question(q_num)

        q_type = _get(r, "QuestionType", "Type", "SavolTuri").lower() or "short_answer"
        prompt = _get(r, "Prompt", "Question", "Text", "Savol", "SavolMatni") or f"Question {q_num}"

        options: list[str] = []
        for letter in ("A", "B", "C", "D", "E", "F", "G", "H"):
            opt_val = _get(r, f"Option{letter}", f"Opt{letter}", letter)
            if opt_val:
                options.append(opt_val)

        if q_type not in QUESTION_TYPES:
            if len(options) >= 2:
                q_type = "multiple_choice"
            else:
                q_type = "short_answer"

        if q_type == "true_false_notgiven" and not options:
            options = ["TRUE", "FALSE", "NOT GIVEN"]
        elif q_type == "yes_no_notgiven" and not options:
            options = ["YES", "NO", "NOT GIVEN"]
        elif q_type in TEXT_TYPES:
            options = []

        ans_str = _get(r, "CorrectAnswer", "Answer", "Key", "Javob", "TogriJavob")
        variants_str = _get(r, "AcceptedVariants", "Variants", "Alternatives", "Variantlar")

        built = build_correct_answers(q_type, options or None, ans_str) if ans_str else []
        if q_type in {"true_false_notgiven", "yes_no_notgiven"} and built:
            correct_answers = [canonical_decision(built[0])]
        elif q_type in CHOICE_TYPES and built:
            correct_answers = built
        else:
            correct_answers = [built[0]] if built else ([ans_str] if ans_str else [])

        accepted_variants = [v.strip() for v in re.split(r"[/;,]+", variants_str) if v.strip()] if variants_str else []
        if len(built) > 1 and q_type in TEXT_TYPES:
            for extra_v in built[1:]:
                if extra_v not in accepted_variants:
                    accepted_variants.append(extra_v)

        if correct_answers:
            with_answers += 1

        pts_raw = _get(r, "Points", "Score", "Ball")
        try:
            points = max(1, int(float(pts_raw))) if pts_raw else 1
        except ValueError:
            points = 1

        wl_raw = _get(r, "WordLimit", "MaxWords", "SozChegarasi")
        try:
            word_limit = int(float(wl_raw)) if wl_raw else None
        except ValueError:
            word_limit = None

        q_item = {
            "number": q_num,
            "partNumber": part_num,
            "type": q_type,
            "prompt": prompt,
            "options": options if options else None,
            "correctAnswers": correct_answers,
            "acceptedVariants": accepted_variants,
            "points": points,
            "wordLimit": word_limit,
            "answerRule": None,
        }
        questions_out.append(q_item)
        by_type[q_type] = by_type.get(q_type, 0) + 1

        p_entry = parts_map.setdefault(
            part_num,
            {
                "partNumber": part_num,
                "title": f"Part {part_num}",
                "passageText": None,
                "instructions": None,
                "questions": [],
            },
        )
        p_entry["questions"].append(q_item)
        passage_val = _get(r, "PassageText", "Passage", "Matn")
        if passage_val and not p_entry["passageText"]:
            p_entry["passageText"] = passage_val
        instr_val = _get(r, "Instructions", "Instruction", "Ko'rsatma", "Korsatma")
        if instr_val and not p_entry["instructions"]:
            p_entry["instructions"] = instr_val

        start_raw = _get(r, "AudioStartSec", "StartSec", "AudioStart")
        end_raw = _get(r, "AudioEndSec", "EndSec", "AudioEnd")
        if start_raw and part_num not in timestamps_map:
            try:
                s_sec = float(start_raw)
                e_sec = float(end_raw) if end_raw else None
                timestamps_map[part_num] = {
                    "partNumber": part_num,
                    "label": f"Part {part_num}",
                    "startSec": s_sec,
                    "endSec": e_sec,
                }
            except ValueError:
                pass

    parts_list = [parts_map[k] for k in sorted(parts_map.keys())]
    timestamps_list = [timestamps_map[k] for k in sorted(timestamps_map.keys())]

    return {
        "parts": parts_list,
        "questions": questions_out,
        "timestamps": timestamps_list,
        "count": len(questions_out),
        "summary": {
            "total": len(questions_out),
            "withAnswers": with_answers,
            "byType": by_type,
        },
    }


def import_spreadsheet_to_exam(
    actor,
    exam_id: str,
    file_bytes: bytes,
    filename: str = "import.xlsx",
    skill: str = "reading",
    group_id: str | None = None,
    replace_existing: bool = True,
) -> dict:
    """Parse an uploaded Excel/CSV file and persist all questions (and timestamps if present) into the exam."""
    parsed = parse_spreadsheet_bytes(file_bytes, filename=filename)
    if group_id:
        res = bulk_import_to_group(
            actor,
            group_id,
            {
                "questions": parsed["questions"],
                "replaceExisting": replace_existing,
                "passageText": next((p["passageText"] for p in parsed["parts"] if p.get("passageText")), None),
                "instructions": next((p["instructions"] for p in parsed["parts"] if p.get("instructions")), None),
            },
        )
    else:
        res = bulk_import_to_exam(
            actor,
            exam_id,
            skill,
            {
                "questions": parsed["questions"],
                "parts": parsed["parts"],
                "replaceExisting": replace_existing,
            },
        )
    if parsed.get("timestamps") and skill == "listening":
        ts_res = save_listening_timestamps(actor, exam_id, parsed["timestamps"])
        res["listeningTimestamps"] = ts_res.get("timestamps")
    res["summary"] = parsed["summary"]
    return res


def save_listening_timestamps(
    actor,
    exam_id: str,
    timestamps: list[dict],
    audio_key: str | None = None,
    audio_url: str | None = None,
    duration_sec: int | None = None,
    uploaded_file=None,
) -> dict:
    """Save single-audio file and section/part timestamp markers for a Listening exam."""
    ensure_listening_config_table()
    try:
        exam = MockExam.objects.get(id=exam_id)
    except MockExam.DoesNotExist:
        raise ContractAPIException("MOCK_EXAM_NOT_FOUND", "Mock imtihon topilmadi", 404)

    _assert_staff_author(actor, exam)

    if uploaded_file is not None:
        audio_key = store_upload(uploaded_file, "audio")

    normalized_ts: list[dict] = []
    if isinstance(timestamps, list):
        for idx, item in enumerate(timestamps):
            if not isinstance(item, dict):
                continue
            part_num = int(item.get("partNumber") or item.get("part") or (idx + 1))
            start_sec = float(item.get("startSec") or item.get("start") or 0.0)
            raw_end = item.get("endSec") if item.get("endSec") is not None else item.get("end")
            end_sec = float(raw_end) if raw_end is not None and str(raw_end).strip() != "" else None
            normalized_ts.append(
                {
                    "partNumber": part_num,
                    "groupId": item.get("groupId"),
                    "label": str(item.get("label") or f"Part {part_num}"),
                    "startSec": max(0.0, round(start_sec, 2)),
                    "endSec": max(0.0, round(end_sec, 2)) if end_sec is not None else None,
                }
            )
    normalized_ts.sort(key=lambda x: x["partNumber"])

    with transaction.atomic():
        section = MockSection.objects.filter(exam_id=exam_id, skill="listening").first()
        groups = (
            list(MockQuestionGroup.objects.filter(section_id=section.id).order_by("sort_order", "id"))
            if section
            else []
        )

        for ts_item in normalized_ts:
            p_num = ts_item["partNumber"]
            matched = next(
                (g for g in groups if g.part_number == p_num or g.sort_order == p_num - 1),
                None,
            )
            if matched:
                ts_item["groupId"] = matched.id
                upd = []
                if audio_key:
                    matched.audio_key = audio_key
                    upd.append("audio_key")
                if ts_item["endSec"] is not None and ts_item["endSec"] > ts_item["startSec"]:
                    part_dur = int(math.ceil(ts_item["endSec"] - ts_item["startSec"]))
                    matched.audio_duration_sec = part_dur
                    upd.append("audio_duration_sec")
                elif duration_sec and not matched.audio_duration_sec:
                    matched.audio_duration_sec = int(duration_sec)
                    upd.append("audio_duration_sec")
                if upd:
                    matched.save(update_fields=upd)

        if audio_key and groups:
            for g in groups:
                if g.audio_key != audio_key:
                    g.audio_key = audio_key
                    g.save(update_fields=["audio_key"])

        resolved_url = audio_url
        if not resolved_url and groups and groups[0].audio_key:
            resolved_url = media_url(groups[0].id, "audio")
        if not audio_key and groups and groups[0].audio_key:
            audio_key = groups[0].audio_key

        if duration_sec is None and normalized_ts:
            max_end = max((t["endSec"] or t["startSec"] for t in normalized_ts), default=0)
            if max_end > 0:
                duration_sec = int(math.ceil(max_end))

        ts_json = json.dumps(normalized_ts)
        section_id = section.id if section else None

        with connection.cursor() as cursor:
            cursor.execute('SELECT 1 FROM "MockListeningConfig" WHERE "exam_id"=%s', [exam_id])
            exists = bool(cursor.fetchone())
            if exists:
                cursor.execute(
                    """
                    UPDATE "MockListeningConfig"
                    SET "section_id"=%s, "audio_key"=COALESCE(%s, "audio_key"),
                        "audio_url"=COALESCE(%s, "audio_url"),
                        "duration_sec"=COALESCE(%s, "duration_sec"),
                        "timestamps"=%s, "updated_at"=CURRENT_TIMESTAMP
                    WHERE "exam_id"=%s
                    """,
                    [section_id, audio_key, resolved_url, duration_sec, ts_json, exam_id],
                )
            else:
                cursor.execute(
                    """
                    INSERT INTO "MockListeningConfig"
                    ("exam_id", "section_id", "audio_key", "audio_url", "duration_sec", "timestamps", "updated_at")
                    VALUES (%s, %s, %s, %s, %s, %s, CURRENT_TIMESTAMP)
                    """,
                    [exam_id, section_id, audio_key, resolved_url, duration_sec, ts_json],
                )

    return get_listening_timestamps(exam_id)


def get_listening_timestamps(exam_id: str) -> dict:
    """Retrieve single-audio configuration and part timestamps for a Listening exam."""
    ensure_listening_config_table()
    with connection.cursor() as cursor:
        cursor.execute(
            """
            SELECT "exam_id", "section_id", "audio_key", "audio_url", "duration_sec", "timestamps"
            FROM "MockListeningConfig"
            WHERE "exam_id"=%s
            """,
            [exam_id],
        )
        row = cursor.fetchone()

    if row:
        ts_val = row[5]
        if isinstance(ts_val, str):
            try:
                ts_val = json.loads(ts_val)
            except Exception:
                ts_val = []
        return {
            "examId": row[0],
            "sectionId": row[1],
            "audioKey": row[2],
            "audioUrl": row[3],
            "durationSec": row[4],
            "timestamps": ts_val if isinstance(ts_val, list) else [],
        }

    # Fallback: synthesize timestamps from existing Listening MockQuestionGroups if present
    section = MockSection.objects.filter(exam_id=exam_id, skill="listening").first()
    if not section:
        return {
            "examId": exam_id,
            "sectionId": None,
            "audioKey": None,
            "audioUrl": None,
            "durationSec": None,
            "timestamps": [],
        }

    groups = list(MockQuestionGroup.objects.filter(section_id=section.id).order_by("sort_order", "id"))
    cursor_sec = 0.0
    fallback_ts = []
    first_audio_key = None
    first_audio_url = None
    for idx, g in enumerate(groups):
        p_num = g.part_number or (idx + 1)
        dur = float(g.audio_duration_sec or 420)
        if g.audio_key and not first_audio_key:
            first_audio_key = g.audio_key
            first_audio_url = media_url(g.id, "audio")
        fallback_ts.append(
            {
                "partNumber": p_num,
                "groupId": g.id,
                "label": g.title or f"Part {p_num}",
                "startSec": round(cursor_sec, 2),
                "endSec": round(cursor_sec + dur, 2),
            }
        )
        cursor_sec += dur

    return {
        "examId": exam_id,
        "sectionId": section.id,
        "audioKey": first_audio_key,
        "audioUrl": first_audio_url,
        "durationSec": int(cursor_sec) if fallback_ts else None,
        "timestamps": fallback_ts,
    }
