"""Group Exam Performance Export Service (Excel & PDF).
Provides styled .xlsx spreadsheets (openpyxl) and official A4 analytical reports (ReportLab)
for dual-layer evaluation results (AI criteria + Teacher final override).
"""
from __future__ import annotations

import os
from datetime import datetime
from io import BytesIO
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from reportlab.lib.colors import HexColor
from reportlab.pdfbase.pdfmetrics import stringWidth
from reportlab.pdfgen import canvas

from apps.legacy_schema.models import (
    AssessmentJob,
    Group,
    MockAttempt,
    MockExam,
    StudentProfile,
    User,
)
from common.api.exceptions import ContractAPIException


def fail(code: str, message: str, status: int = 400):
    raise ContractAPIException(code=code, message=message, status=status)


def compute_cefr_level(score: Optional[float], existing_cefr: Optional[str] = None) -> str:
    """Resolve CEFR level from existing value or calculated band / percentage score."""
    if existing_cefr and str(existing_cefr).strip():
        return str(existing_cefr).strip().upper()
    if score is None:
        return "-"
    try:
        val = float(score)
    except (ValueError, TypeError):
        return "-"
    # Multi-level standard scale (0 - 75 points) or IELTS band (0 - 9.0)
    if val >= 70 or (val <= 9.0 and val >= 7.0):
        return "C1"
    if val >= 55 or (val <= 9.0 and val >= 5.5):
        return "B2"
    if val >= 40 or (val <= 9.0 and val >= 4.0):
        return "B1"
    if val >= 30 or (val <= 9.0 and val >= 3.0):
        return "A2"
    return "A1"


def get_group_exam_export_data(actor: User, group_id: str, exam_id: str) -> Dict[str, Any]:
    """Gather complete group roster, attempts, AI criteria, and teacher scores for export."""
    if getattr(actor, "role", None) not in ("teacher", "admin", "super_admin"):
        fail("FORBIDDEN", "Ushbu amal uchun ruxsat berilmagan", 403)

    group = Group.objects.filter(id=group_id).select_related("teacher").first()
    if group is None:
        fail("GROUP_NOT_FOUND", "Guruh topilmadi", 404)

    if actor.role == "teacher" and group.teacher_id != actor.id:
        fail("FORBIDDEN", "Bu guruh sizga biriktirilmagan", 403)

    exam = MockExam.objects.filter(id=exam_id).first()
    if exam is None:
        fail("EXAM_NOT_FOUND", "Imtihon topilmadi", 404)

    # 1. Group student roster
    roster = list(
        StudentProfile.objects.filter(group_id=group.id)
        .select_related("user")
        .order_by("user__name", "user_id")
    )
    student_ids = [student.user_id for student in roster]

    # 2. Latest attempt per student for this exam (in-memory grouping for SQLite/PostgreSQL compatibility)
    attempts_qs = (
        MockAttempt.objects.filter(exam_id=exam.id, student_id__in=student_ids)
        .order_by("-submitted_at", "-started_at", "-id")
    )
    attempts_by_student: Dict[str, MockAttempt] = {}
    for att in attempts_qs:
        if att.student_id not in attempts_by_student:
            attempts_by_student[att.student_id] = att

    attempt_ids = [att.id for att in attempts_by_student.values()]

    # 3. Assessment jobs (AI evaluations)
    jobs = list(
        AssessmentJob.objects.filter(attempt_id__in=attempt_ids)
        .order_by("-created_at", "-id")
    )
    jobs_by_attempt_skill: Dict[tuple, AssessmentJob] = {}
    for job in jobs:
        key = (job.attempt_id, job.skill)
        if key not in jobs_by_attempt_skill:
            jobs_by_attempt_skill[key] = job

    # 4. Build student rows
    students_data: List[Dict[str, Any]] = []
    submitted_count = 0
    final_scores_list: List[float] = []
    writing_scores_list: List[float] = []
    speaking_scores_list: List[float] = []
    cefr_counts: Dict[str, int] = {"C1": 0, "B2": 0, "B1": 0, "A2": 0, "A1": 0}

    for idx, st in enumerate(roster, start=1):
        att = attempts_by_student.get(st.user_id)
        u = getattr(st, "user", None)
        st_name = getattr(u, "name", "Noma'lum") or "Noma'lum"
        st_phone = getattr(u, "phone", "-") or "-"

        if att:
            submitted_count += 1
            sec_bands = dict(att.section_bands or {})
            std_scores = dict(att.standard_scores or {})
            media_st = dict(att.media_state or {})

            w_job = jobs_by_attempt_skill.get((att.id, "writing"))
            s_job = jobs_by_attempt_skill.get((att.id, "speaking"))

            # Writing score
            w_score = None
            if "writing" in sec_bands and sec_bands["writing"] is not None:
                w_score = float(sec_bands["writing"])
            elif "writing" in std_scores and std_scores["writing"] is not None:
                w_score = float(std_scores["writing"])
            elif w_job and w_job.final_score is not None:
                w_score = float(w_job.final_score)
            elif w_job and w_job.ai_score is not None:
                w_score = float(w_job.ai_score)

            # Speaking score
            s_score = None
            if "speaking" in sec_bands and sec_bands["speaking"] is not None:
                s_score = float(sec_bands["speaking"])
            elif "speaking" in std_scores and std_scores["speaking"] is not None:
                s_score = float(std_scores["speaking"])
            elif s_job and s_job.final_score is not None:
                s_score = float(s_job.final_score)
            elif s_job and s_job.ai_score is not None:
                s_score = float(s_job.ai_score)

            if w_score is not None:
                writing_scores_list.append(w_score)
            if s_score is not None:
                speaking_scores_list.append(s_score)

            # AI score calculation (average of AI jobs if available)
            ai_parts = []
            if w_job and w_job.ai_score is not None:
                ai_parts.append(float(w_job.ai_score))
            if s_job and s_job.ai_score is not None:
                ai_parts.append(float(s_job.ai_score))
            ai_score = round(sum(ai_parts) / len(ai_parts), 1) if ai_parts else None

            # Teacher score
            teacher_score = None
            if media_st.get("teacher_score") is not None:
                teacher_score = float(media_st["teacher_score"])
            elif w_job and w_job.teacher_score is not None:
                teacher_score = float(w_job.teacher_score)
            elif s_job and s_job.teacher_score is not None:
                teacher_score = float(s_job.teacher_score)

            # Final score
            final_score = att.overall_band if att.overall_band is not None else att.overall_score
            if final_score is not None:
                final_score = float(final_score)
                final_scores_list.append(final_score)

            # CEFR Level
            cefr = compute_cefr_level(final_score, att.cefr_level)
            if cefr in cefr_counts:
                cefr_counts[cefr] += 1

            # Status determination
            is_teacher_reviewed = bool(
                media_st.get("is_teacher_reviewed")
                or media_st.get("is_reviewed_by_teacher")
                or (w_job and w_job.final_score_source == "TEACHER")
                or (s_job and s_job.final_score_source == "TEACHER")
            )
            if is_teacher_reviewed:
                status_label = "Ko'rib chiqildi"
                status_code = "REVIEWED"
            elif att.status == "completed":
                status_label = "Baholandi (AI)"
                status_code = "COMPLETED"
            elif att.status == "in_progress":
                status_label = "Jarayonda"
                status_code = "IN_PROGRESS"
            else:
                status_label = "Topshirilgan"
                status_code = "SUBMITTED"

            date_str = (
                att.submitted_at.strftime("%Y-%m-%d %H:%M")
                if att.submitted_at
                else (att.started_at.strftime("%Y-%m-%d %H:%M") if att.started_at else "-")
            )

            students_data.append({
                "number": idx,
                "student_id": st.user_id,
                "name": st_name,
                "phone": st_phone,
                "attempt_id": att.id,
                "submitted_at": date_str,
                "writing_score": w_score,
                "speaking_score": s_score,
                "ai_score": ai_score,
                "teacher_score": teacher_score,
                "final_score": final_score,
                "cefr_level": cefr,
                "status": status_label,
                "status_code": status_code,
                "is_teacher_reviewed": is_teacher_reviewed,
                "teacher_comment": media_st.get("teacher_comment", ""),
            })
        else:
            students_data.append({
                "number": idx,
                "student_id": st.user_id,
                "name": st_name,
                "phone": st_phone,
                "attempt_id": None,
                "submitted_at": "Topshirilmagan",
                "writing_score": None,
                "speaking_score": None,
                "ai_score": None,
                "teacher_score": None,
                "final_score": None,
                "cefr_level": "-",
                "status": "Topshirilmagan",
                "status_code": "NOT_ATTEMPTED",
                "is_teacher_reviewed": False,
                "teacher_comment": "",
            })

    # Summary metrics
    avg_final = round(sum(final_scores_list) / len(final_scores_list), 1) if final_scores_list else 0.0
    avg_writing = round(sum(writing_scores_list) / len(writing_scores_list), 1) if writing_scores_list else 0.0
    avg_speaking = round(sum(speaking_scores_list) / len(speaking_scores_list), 1) if speaking_scores_list else 0.0

    teacher_name = "Noma'lum"
    if group.teacher:
        teacher_name = group.teacher.name or group.teacher.phone

    return {
        "group": {
            "id": group.id,
            "name": group.name,
            "teacherName": teacher_name,
        },
        "exam": {
            "id": exam.id,
            "title": exam.title,
            "type": exam.type,
        },
        "totalStudents": len(roster),
        "submittedCount": submitted_count,
        "submissionRate": round((submitted_count / len(roster) * 100), 1) if roster else 0.0,
        "averageFinalScore": avg_final,
        "averageWritingScore": avg_writing,
        "averageSpeakingScore": avg_speaking,
        "cefrDistribution": cefr_counts,
        "students": students_data,
        "exportedAt": datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
    }


def generate_group_exam_excel(data: Dict[str, Any]) -> bytes:
    """Generate professional, emerald-themed .xlsx spreadsheet."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Imtihon Natijalari"

    # Color Palette
    primary_emerald = "064E3B"   # Deep Dark Emerald
    accent_emerald = "059669"    # Emerald 600
    light_emerald = "D1FAE5"     # Emerald 100
    white = "FFFFFF"
    zebra_light = "F8FAFC"       # Light Slate
    border_gray = "CBD5E1"
    text_dark = "0F172A"

    font_title = Font(name="Calibri", size=16, bold=True, color=primary_emerald)
    font_subtitle = Font(name="Calibri", size=11, bold=True, color=accent_emerald)
    font_meta_bold = Font(name="Calibri", size=10, bold=True, color=text_dark)
    font_meta = Font(name="Calibri", size=10, color="475569")
    font_header = Font(name="Calibri", size=10, bold=True, color=white)
    font_data = Font(name="Calibri", size=10, color=text_dark)
    font_data_bold = Font(name="Calibri", size=10, bold=True, color=text_dark)

    fill_header = PatternFill(start_color=primary_emerald, end_color=primary_emerald, fill_type="solid")
    fill_subtotal = PatternFill(start_color=light_emerald, end_color=light_emerald, fill_type="solid")
    fill_zebra = PatternFill(start_color=zebra_light, end_color=zebra_light, fill_type="solid")
    fill_white = PatternFill(start_color=white, end_color=white, fill_type="solid")

    thin_border = Side(border_style="thin", color=border_gray)
    cell_border = Border(left=thin_border, right=thin_border, top=thin_border, bottom=thin_border)

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")
    align_right = Alignment(horizontal="right", vertical="center")

    # 1. Title Block
    ws.merge_cells("A1:K1")
    ws["A1"] = os.environ.get("CENTER_NAME", "GRAND EDUCATION & BESTWAY ACADEMY")
    ws["A1"].font = font_title
    ws["A1"].alignment = align_center

    ws.merge_cells("A2:K2")
    ws["A2"] = "IMTIHON NATIJALARI VA AKADEMIK TAHLIL HISOBOTI (DUAL-LAYER GRADING)"
    ws["A2"].font = font_subtitle
    ws["A2"].alignment = align_center

    # 2. Meta Info Card (Rows 4-5)
    grp = data["group"]
    exm = data["exam"]

    ws["A4"] = "Guruh:"
    ws["A4"].font = font_meta_bold
    ws["B4"] = grp["name"]
    ws["B4"].font = font_meta

    ws["D4"] = "Imtihon:"
    ws["D4"].font = font_meta_bold
    ws["E4"] = exm["title"]
    ws["E4"].font = font_meta

    ws["H4"] = "O'qituvchi:"
    ws["H4"].font = font_meta_bold
    ws["I4"] = grp["teacherName"]
    ws["I4"].font = font_meta

    ws["A5"] = "Jami O'quvchilar:"
    ws["A5"].font = font_meta_bold
    ws["B5"] = f"{data['totalStudents']} nafar (Topshirdi: {data['submittedCount']})"
    ws["B5"].font = font_meta

    ws["D5"] = "O'rtacha Yakuniy Ball:"
    ws["D5"].font = font_meta_bold
    ws["E5"] = str(data["averageFinalScore"])
    ws["E5"].font = font_meta

    ws["H5"] = "Eksport Sanasi:"
    ws["H5"].font = font_meta_bold
    ws["I5"] = data["exportedAt"]
    ws["I5"].font = font_meta

    # 3. Table Headers (Row 7)
    headers = [
        ("№", 5, align_center),
        ("O'quvchi Ismi", 28, align_left),
        ("Telefon", 16, align_center),
        ("Topshirilgan Sana", 18, align_center),
        ("Writing", 10, align_center),
        ("Speaking", 10, align_center),
        ("AI Balli", 10, align_center),
        ("Ustoz Balli", 12, align_center),
        ("Yakuniy Ball", 13, align_center),
        ("CEFR", 9, align_center),
        ("Holat", 18, align_center),
    ]

    header_row = 7
    for col_idx, (col_name, col_width, _) in enumerate(headers, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=col_name)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = align_center
        cell.border = cell_border
        col_letter = openpyxl.utils.get_column_letter(col_idx)
        ws.column_dimensions[col_letter].width = col_width

    # 4. Data Rows
    current_row = header_row + 1
    for st in data["students"]:
        row_fill = fill_zebra if current_row % 2 == 0 else fill_white
        
        def fmt_num(v):
            return f"{v:.1f}" if isinstance(v, (int, float)) else ("-" if v is None else str(v))

        row_values = [
            (st["number"], align_center, font_data),
            (st["name"], align_left, font_data_bold),
            (st["phone"], align_center, font_data),
            (st["submitted_at"], align_center, font_data),
            (fmt_num(st["writing_score"]), align_center, font_data),
            (fmt_num(st["speaking_score"]), align_center, font_data),
            (fmt_num(st["ai_score"]), align_center, font_data),
            (fmt_num(st["teacher_score"]), align_center, font_data_bold),
            (fmt_num(st["final_score"]), align_center, Font(name="Calibri", size=10, bold=True, color=primary_emerald)),
            (st["cefr_level"], align_center, Font(name="Calibri", size=10, bold=True, color=accent_emerald)),
            (st["status"], align_center, font_data),
        ]

        for col_idx, (val, alignment, font_style) in enumerate(row_values, start=1):
            cell = ws.cell(row=current_row, column=col_idx, value=val)
            cell.font = font_style
            cell.fill = row_fill
            cell.alignment = alignment
            cell.border = cell_border

        current_row += 1

    # 5. Summary / Subtotals
    current_row += 1
    ws.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=4)
    summary_label = ws.cell(row=current_row, column=1, value="Guruh Bo'yicha O'rtacha Ko'rsatkichlar:")
    summary_label.font = font_meta_bold
    summary_label.alignment = align_right
    summary_label.fill = fill_subtotal

    for c in range(1, 5):
        ws.cell(row=current_row, column=c).fill = fill_subtotal
        ws.cell(row=current_row, column=c).border = cell_border

    ws.cell(row=current_row, column=5, value=f"{data['averageWritingScore']:.1f}").alignment = align_center
    ws.cell(row=current_row, column=6, value=f"{data['averageSpeakingScore']:.1f}").alignment = align_center
    ws.cell(row=current_row, column=7, value="-").alignment = align_center
    ws.cell(row=current_row, column=8, value="-").alignment = align_center
    ws.cell(row=current_row, column=9, value=f"{data['averageFinalScore']:.1f}").alignment = align_center
    ws.cell(row=current_row, column=10, value="-").alignment = align_center
    ws.cell(row=current_row, column=11, value=f"{data['submissionRate']}% topshirdi").alignment = align_center

    for c in range(5, 12):
        ws.cell(row=current_row, column=c).font = font_data_bold
        ws.cell(row=current_row, column=c).fill = fill_subtotal
        ws.cell(row=current_row, column=c).border = cell_border

    # 6. CEFR Breakdown block
    current_row += 2
    ws.cell(row=current_row, column=1, value="CEFR Taqsimoti:").font = font_meta_bold
    cefr_dist = data.get("cefrDistribution", {})
    dist_str = f"C1: {cefr_dist.get('C1', 0)} ta   |   B2: {cefr_dist.get('B2', 0)} ta   |   B1: {cefr_dist.get('B1', 0)} ta   |   A2: {cefr_dist.get('A2', 0)} ta"
    ws.merge_cells(start_row=current_row, start_column=2, end_row=current_row, end_column=6)
    cefr_cell = ws.cell(row=current_row, column=2, value=dist_str)
    cefr_cell.font = font_meta

    output = BytesIO()
    wb.save(output)
    return output.getvalue()


def generate_group_exam_pdf(data: Dict[str, Any]) -> bytes:
    """Generate official A4 analytical report PDF using ReportLab."""
    output = BytesIO()
    width, height = 595.28, 841.89  # A4 Portrait
    doc = canvas.Canvas(output, pagesize=(width, height))

    primary = "#064e3b"   # Dark Emerald
    accent = "#059669"    # Emerald 600
    slate_dark = "#0f172a"
    slate_muted = "#475569"
    border_color = "#cbd5e1"

    # Outer Decorative Double Frame
    doc.setStrokeColor(HexColor(primary))
    doc.setLineWidth(2.5)
    doc.rect(26, 26, width - 52, height - 52)

    doc.setStrokeColor(HexColor(accent))
    doc.setLineWidth(0.8)
    doc.rect(32, 32, width - 64, height - 64)

    def draw_text(val, x, y, size=10, font="Helvetica", color=slate_dark, center=False):
        val = str(val)
        doc.setFillColor(HexColor(color))
        if center:
            x = (width - stringWidth(val, font, size)) / 2
        item = doc.beginText(x, height - y)
        item.setFont(font, size)
        item.textOut(val)
        doc.drawText(item)

    def center_text(val, y, size=11, bold=False, color=slate_muted):
        font = "Helvetica-Bold" if bold else "Helvetica"
        draw_text(val, 0, y, size=size, font=font, color=color, center=True)

    # 1. Header Branding
    center_text(os.environ.get("CENTER_NAME", "GRAND EDUCATION & BESTWAY ACADEMY"), 60, 18, True, primary)
    center_text("XALQARO IMTIHONLARGA TAYYORLASH VA DUAL-LAYER BAHOLASH TIZIMI", 78, 8, False, "#64748b")
    center_text("IMTIHON NATIJALARI VA AKADEMIK TAHLIL RASMIY HISOBOTI", 96, 11, True, accent)

    # 2. Meta Info Banner
    grp = data["group"]
    exm = data["exam"]

    doc.setFillColor(HexColor("#f8fafc"))
    doc.rect(44, height - 160, width - 88, 48, fill=1, stroke=0)
    doc.setStrokeColor(HexColor(border_color))
    doc.setLineWidth(0.6)
    doc.rect(44, height - 160, width - 88, 48, fill=0, stroke=1)

    draw_text(f"Guruh: {grp['name']}", 54, 126, 9.5, "Helvetica-Bold", slate_dark)
    draw_text(f"Ustoz: {grp['teacherName']}", 54, 142, 8.5, "Helvetica", slate_muted)

    draw_text(f"Imtihon: {exm['title'][:32]}", 230, 126, 9.5, "Helvetica-Bold", slate_dark)
    draw_text(f"O'quvchilar: {data['totalStudents']} nafar (Topshirdi: {data['submittedCount']})", 230, 142, 8.5, "Helvetica", slate_muted)

    draw_text(f"O'rtacha Ball: {data['averageFinalScore']}", 430, 126, 10, "Helvetica-Bold", accent)
    draw_text(f"Sana: {data['exportedAt'][:10]}", 430, 142, 8.5, "Helvetica", slate_muted)

    # 3. Table Headers
    y = 180
    headers = [
        ("№", 48),
        ("Talaba F.I.Sh", 68),
        ("Telefon", 185),
        ("Writing", 260),
        ("Speaking", 310),
        ("AI", 365),
        ("Ustoz", 400),
        ("Yakuniy", 440),
        ("CEFR", 485),
        ("Holat", 518),
    ]

    doc.setFillColor(HexColor("#f1f5f9"))
    doc.rect(44, height - y - 13, width - 88, 17, fill=1, stroke=0)
    for title, x_pos in headers:
        draw_text(title, x_pos, y, 8, "Helvetica-Bold", slate_dark)

    y += 17
    doc.setStrokeColor(HexColor(border_color))
    doc.line(44, height - y + 4, width - 44, height - y + 4)

    # 4. Student Rows (Render up to 20 rows on A4)
    students = data.get("students", [])[:20]
    for idx, st in enumerate(students):
        row_bg = "#ffffff" if idx % 2 == 0 else "#f8fafc"
        doc.setFillColor(HexColor(row_bg))
        doc.rect(44, height - y - 10, width - 88, 15, fill=1, stroke=0)

        def fv(v):
            return f"{v:.1f}" if isinstance(v, (int, float)) else ("-" if v is None else str(v))

        draw_text(str(st["number"]), 48, y, 7.5, "Helvetica", slate_muted)
        draw_text(st["name"][:20], 68, y, 7.5, "Helvetica-Bold", slate_dark)
        draw_text(st["phone"][:12], 185, y, 7, "Helvetica", slate_muted)
        draw_text(fv(st["writing_score"]), 265, y, 7.5, "Helvetica", slate_dark)
        draw_text(fv(st["speaking_score"]), 315, y, 7.5, "Helvetica", slate_dark)
        draw_text(fv(st["ai_score"]), 368, y, 7.5, "Helvetica", slate_muted)
        draw_text(fv(st["teacher_score"]), 403, y, 7.5, "Helvetica-Bold", accent)
        draw_text(fv(st["final_score"]), 443, y, 8, "Helvetica-Bold", primary)
        draw_text(st["cefr_level"], 488, y, 8, "Helvetica-Bold", accent)

        st_label = st["status"][:10]
        st_color = "#16a34a" if "Tekshirildi" in st["status"] or "Ko'rib" in st["status"] else (
            "#2563eb" if "Baholandi" in st["status"] else "#94a3b8"
        )
        draw_text(st_label, 518, y, 7, "Helvetica-Bold", st_color)

        y += 15
        if y > 600:
            break

    doc.setStrokeColor(HexColor(border_color))
    doc.line(44, height - y + 4, width - 44, height - y + 4)

    # 5. CEFR Distribution Breakdown Box
    y += 18
    draw_text("CEFR DARAJALARI BO'YICHA GURUH TAQSIMOTI", 44, y, 9, "Helvetica-Bold", primary)
    y += 14

    cd = data.get("cefrDistribution", {})
    cefr_summary_text = (
        f"• C1 (70+ ball): {cd.get('C1', 0)} nafar    "
        f"• B2 (55-69 ball): {cd.get('B2', 0)} nafar    "
        f"• B1 (40-54 ball): {cd.get('B1', 0)} nafar    "
        f"• A2: {cd.get('A2', 0)} nafar"
    )
    draw_text(cefr_summary_text, 50, y, 8, "Helvetica", slate_dark)

    y += 14
    stats_summary = (
        f"O'rtacha Writing: {data['averageWritingScore']}  |  "
        f"O'rtacha Speaking: {data['averageSpeakingScore']}  |  "
        f"Guruh O'rtacha Yakuniy Balli: {data['averageFinalScore']}  |  "
        f"Topshirish ko'rsatkichi: {data['submissionRate']}%"
    )
    draw_text(stats_summary, 50, y, 8, "Helvetica-Bold", accent)

    # 6. Official Signatures and Seal
    sig_y = height - 110
    doc.setStrokeColor(HexColor("#94a3b8"))
    doc.line(54, sig_y, 210, sig_y)
    doc.line(width - 210, sig_y, width - 54, sig_y)

    item1 = doc.beginText(54, sig_y - 14)
    item1.setFont("Helvetica", 8)
    item1.textOut("Guruh O'qituvchisi Imzosi")
    doc.drawText(item1)

    item2 = doc.beginText(width - 210, sig_y - 14)
    item2.setFont("Helvetica", 8)
    item2.textOut("Akademik Bo'lim Rahbari / Muhr O'rni")
    doc.drawText(item2)

    # 7. Document Verification Footer
    verify_url = f"https://bestwayec.uz/groups/{grp['id']}"
    draw_text(
        f"Hujjat haqiqiyligi tekshiruvi: {verify_url} · Grand Education Dual-Layer Assessment System",
        44,
        height - 40,
        7,
        color="#94a3b8",
    )

    doc.showPage()
    doc.save()
    return output.getvalue()
